"""
PAINEL DE FUNDOS - um arquivo so (dados abertos CVM + ANBIMA + Banco Central).

Como usar em qualquer computador:
  1) Instale o Python 3.11+ (python.org; no Windows marque "Add python.exe to PATH").
  2) No terminal:   pip install fastapi uvicorn pandas numpy requests pyarrow openpyxl
  3) Salve este arquivo como painel.py e rode:   python painel.py
     -> baixa e processa NOSSOS_FUNDOS + PEERS (listas abaixo, desde DESDE) e so depois abre o painel
        no navegador em http://localhost:7860 com tudo pronto.
  Excel em vez do painel:   python painel.py excel

Fontes: CVM (CDA mensal, informe diario, cadastro), ANBIMA (debentures), Banco Central (CDI). Tudo publico e gratuito.
Os dados ficam na pasta dados_painel/, ao lado deste arquivo (a 1a vez baixa ~2-3 GB da CVM; evite pasta do OneDrive).
Depois so baixa o que mudou; o painel processado do dia e reaproveitado (abre em segundos).
Arquivos mensais republicados pela CVM (fundos tem ate 3 meses para abrir a carteira toda) sao rebaixados 1x por dia.
"""
import sys
import io
import json
import os
import re
import time
import zipfile
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import requests

# ------------------------------------------------------------------ CONFIGURACAO (edite aqui)
NOSSOS_FUNDOS = {               # nome curto: CNPJ (qualquer formato)
    "DUAL": "34.803.938/0001-61",
    "DUAL_GLOBAL": "62.917.953/0001-76",
    "PRECISION_PM": "32.292.528/0001-78",
}
PEERS = {
    "AZ_ALTRO": "22.100.009/0001-07", "SPARTA_TOP": "14.188.162/0001-00", "XP_CE120": "22.003.930/0001-31",
    "CAPITANIA_P45": "20.146.294/0001-71", "KINEA_CP_PREV": "26.491.419/0001-87", "IBIUNA_CREDIT": "37.310.657/0001-65",
    "DAYCOVAL_CLASSIC": "10.783.480/0001-68", "RIZA_LOTUS_PREV": "43.423.186/0001-02",
    "BRADESCO_CP_PLUS": "32.387.924/0001-89", "MAPFRE_CONFIANZA": "51.253.495/0001-00", "CAIXA_MAXI": "17.322.725/0001-07",
    "KINEA_RF_CP": "41.978.506/0001-57", "REGIA_EQUILIBRIO": "53.828.295/0001-55", "BNP_CREDITO_PLUS": "17.137.984/0001-50",
    "OCCAM_LIQUIDEZ": "46.098.897/0001-39", "SANTANDER_INFRA_CDI": "51.672.063/0001-25", "BTG_CRED_CORP": "14.557.317/0001-38",
    "XP_LIQUIDEZ": "51.488.342/0001-33", "SPX_SEAHAWK": "35.491.217/0001-26", "WESTERN_TOTAL_CREDIT": "28.320.756/0001-37",
    "SULAMERICA_CRED_ATIVO": "13.823.084/0001-05", "SAFRA_VITESSE": "58.735.449/0001-88", "JGP_DEB_CDI": "58.600.298/0001-50",
    "COMPASS_CREDIT": "35.399.404/0001-84", "SVN_RF_CP": "51.825.326/0001-99", "VINLAND_CORE": "56.415.717/0001-59",
    "POLO_CRED_CORP": "56.974.598/0001-74", "ASA_ALM": "50.911.242/0001-05",
    "SICOOB_INSTITUCIONAL": "14.702.111/0001-54", "ICATU_VANGUARDA": "64.203.379/0001-10", "SOMMA_QP": "24.249.979/0001-02",
    "ANGA_CRED_ESTR": "23.034.819/0001-75", "MAG_ZONA_MATA": "41.594.651/0001-34", "LEGACY_COMPOUND": "50.891.130/0001-30",
    "XP_AUGME_XPCE": "67.007.466/0001-90", "ARX_INFRA": "63.920.768/0001-01", "UBS_EVOLUTION": "56.049.361/0001-87",
    "WRIGHT_CRED2": "53.179.441/0001-69", "VALORA_ABSOLUTE": "10.326.625/0001-00", "INTER_POLARIS": "64.156.687/0001-31",
    "SICREDI_INFRA": "61.734.698/0001-63", "PLURAL_DEB_INC": "58.052.836/0001-10", "CAPITANIA_INFRA90": "52.248.139/0001-52",
    "AUGME_MRT2": "27.347.344/0001-28", "V8_MERCURY": "58.398.452/0001-53", "A1_HIGH_GRADE": "57.815.131/0001-44",
    "BANRISUL_CABERGS": "05.196.208/0001-41", "KILIMA_BANCOS": "49.272.086/0001-09", "PRINZ_LIQUIDEZ": "59.376.795/0001-80",
    "TENAX_RFA": "53.293.548/0001-33", "DRYS_SHELTER": "52.282.978/0001-97", "JOURNEY_JCW": "57.594.567/0001-50",
}
DESDE = "2019-01"               # primeiro mes de carteira/cota (CDA existe desde 2005; mais antigo = mais download na 1a vez)
# grupo economico: a CVM so informa o emissor. Regra = trecho do nome do emissor (sem acento, maiusculo) -> grupo.
# Securitizadoras (Opea, Vert, True...) ficam como estao: o devedor do lastro nao vem na CDA.
GRUPOS = {
    "LOCALIZA|UNIDAS": "Localiza", "REDE D.?OR|DOR SAO LUIZ": "Rede D'Or", "AEGEA|AGUAS DO RIO|AGUAS DO PARA|EQUIPAV|CORSAN|RIOGRANDENSE DE SANEAMENTO": "Aegea",
    "BRK AMBIENTAL": "BRK", "EQUATORIAL|CEEE": "Equatorial", "ENERGISA": "Energisa", "CEMIG": "Cemig", "COPEL": "Copel",
    "SABESP|SANEAMENTO BASICO DO ESTADO DE SAO PAULO": "Sabesp", "COPASA|SANEAMENTO DE MINAS": "Copasa", "PETROBRAS|PETROLEO BRASILEIRO": "Petrobras",
    "PRIO|PETRO RIO": "Prio", "PETRORECONCAVO": "PetroReconcavo", "COSAN|RAIZEN|RUMO|COMPASS|MOOVE": "Cosan", "VALE S": "Vale",
    "SUZANO": "Suzano", "KLABIN": "Klabin", "SENDAS|ASSAI": "Assaí", "HAPVIDA|NOTRE DAME": "Hapvida", "VAMOS|SIMPAR|JSL|MOVIDA|CS INFRA": "Simpar",
    "MILLS": "Mills", "EUROFARMA": "Eurofarma", "IOCHPE": "Iochpe-Maxion", "CCR|MOTIVA": "Motiva (CCR)", "ECORODOVIAS": "EcoRodovias",
    "ELETROBRAS|CENTRAIS ELETRICAS BRASILEIRAS|FURNAS|CHESF|ELETRONORTE": "Eletrobras", "TAESA|TRANSMISSORA ALIANCA": "Taesa",
    "ISA ENERGIA|CTEEP": "ISA Energia", "AUREN|CESP|AES BRASIL": "Auren", "NOVA TRANSPORTADORA DO SUDESTE": "NTS", "VLI ": "VLI",
    "ITAU": "Itaú", "BRADESCO": "Bradesco", "SANTANDER": "Santander", "BANCO DO BRASIL": "Banco do Brasil", "CAIXA ECONOMICA": "Caixa Econômica",
    "BTG": "BTG Pactual", "VOTORANTIM|BANCO BV": "BV", "XP ": "XP", "SAFRA": "Safra", "DAYCOVAL": "Daycoval",
    "TESOURO NACIONAL|SECRETARIA DO TESOURO": "Tesouro Nacional",
}
OUT = "lookthrough_cvm.xlsx"
WORKERS = 4

URL = "https://dados.cvm.gov.br/dados/FI/DOC"
CACHE = Path(os.environ.get("PAINEL_DADOS") or Path(__file__).resolve().parent / "dados_painel")   # pasta ao lado do .py
(CACHE / "rec").mkdir(parents=True, exist_ok=True)
VERSAO = 2                      # muda quando o recorte guarda colunas novas (forca reprocessar)
REN = {"CNPJ_FUNDO": "CNPJ_FUNDO_CLASSE", "CNPJ_FUNDO_COTA": "CNPJ_FUNDO_CLASSE_COTA",
       "NM_FUNDO_COTA": "NM_FUNDO_CLASSE_SUBCLASSE_COTA"}             # nomes antigos (ate 2022) -> novos
NUM = ["QT_POS_FINAL", "VL_CUSTO_POS_FINAL", "VL_AQUIS_NEGOC", "VL_VENDA_NEGOC", "PR_INDEXADOR_POSFX",
       "PR_CUPOM_POSFX", "PR_TAXA_PREFX"]
COLS = {"CNPJ_FUNDO_CLASSE", "DT_COMPTC", "TP_APLIC", "TP_ATIVO", "VL_MERC_POS_FINAL", "CNPJ_FUNDO_CLASSE_COTA",
        "NM_FUNDO_CLASSE_SUBCLASSE_COTA", "DS_ATIVO", "CD_ATIVO", "CD_ISIN", "TP_TITPUB", "DT_VENC", "EMISSOR",
        "CNPJ_EMISSOR", "CPF_CNPJ_EMISSOR", "DS_SWAP", "CD_SWAP", "DS_ATIVO_EXTERIOR", "CD_ATIVO_BV_MERC", "PAIS",
        "DS_INDEXADOR_POSFX", "GRAU_RISCO", "AG_RISCO", *NUM} | set(REN)
PASSIVO = r"(?i)a pagar|obriga|exigibilidade"                        # entram negativos (PL = ativos - passivos)
DERIV = r"(?i)mercado futuro|opç|swap|termo"   # derivativos vem pelo nocional: fora da soma de 100% e dos graficos
AJUSTE = "Ajuste carteira x PL"


def fmt(c):
    d = re.sub(r"\D", "", c)
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"


def cnpj_of(x):
    if "maisretorno" in x:
        m = re.search(r"id=(\d{14})", x) or re.search(
            r"CNPJ\D{0,30}(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})", requests.get(x, timeout=60, headers={"User-Agent": "Mozilla/5.0"}).text)
        x = m.group(1)
    return fmt(x)


def baixa(url):
    f = CACHE / url.rsplit("/", 1)[1]
    if not f.exists():
        tmp = f.with_suffix(f".{os.getpid()}.part")
        with requests.get(url, stream=True, timeout=600) as r:
            r.raise_for_status()
            with open(tmp, "wb") as o:
                for ch in r.iter_content(1 << 20):
                    o.write(ch)
        tmp.replace(f)
    return f


def lista(pasta, padrao):        # arquivos disponiveis no diretorio da CVM
    return sorted(set(re.findall(padrao, requests.get(f"{URL}/{pasta}", timeout=60).text)))


def arquivos(tipo, desde):       # {arquivo: [meses]} - mensal recente + anual (HIST) antigo
    mens = lista(f"{tipo}/DADOS/", rf"{tipo.lower()}_fi_(\d{{6}})\.zip")
    anos = lista(f"{tipo}/DADOS/HIST/", rf"{tipo.lower()}_fi_(\d{{4}})\.zip")
    out = {}
    for a in anos:
        ms = [f"{a}-{m:02d}" for m in range(1, 13) if f"{a}-{m:02d}" >= desde and f"{a}{m:02d}" not in mens]
        if ms:
            out[f"{URL}/{tipo}/DADOS/HIST/{tipo.lower()}_fi_{a}.zip"] = ms
    for m in mens:
        if f"{m[:4]}-{m[4:]}" >= desde:
            out[f"{URL}/{tipo}/DADOS/{tipo.lower()}_fi_{m}.zip"] = [f"{m[:4]}-{m[4:]}"]
    return out


def le(z, nome, cols):
    df = pd.read_csv(z.open(nome), sep=";", encoding="latin1", dtype=str, low_memory=False, quoting=3,
                     usecols=lambda c: c in cols)                      # quoting=3: CVM nao usa aspas (e tem " soltas)
    return df.rename(columns=REN)


def marca(kind, m):
    f = CACHE / "rec" / f"{kind}_{m}.json"
    try:
        j = json.loads(f.read_text())
        return set(j["fundos"]) if isinstance(j, dict) and j.get("v") == VERSAO else None
    except (OSError, ValueError):
        return None


def fecho(lk, peers, base=()):   # tudo que os peers alcancam via cotas de fundos
    alvo, novo = set(peers) | set(base), set(peers)
    while novo:
        novo = set(lk[lk.CNPJ_FUNDO_CLASSE.isin(novo)].CNPJ_FUNDO_CLASSE_COTA) - alvo
        alvo |= novo
    return alvo


def processa_cda(url, meses, peers):
    """Le um zip CDA: links fundo->cota e PL de todos; posicoes so dos fundos ligados aos peers."""
    z = zipfile.ZipFile(baixa(url))
    nomes = z.namelist()
    if not any("_PL_" in n for n in nomes) or not any("_BLC_2_" in n for n in nomes):
        return url + " (CVM ainda publicando este mes: pulado)"
    pl = le(z, next(n for n in nomes if "_PL_" in n), {"CNPJ_FUNDO_CLASSE", "CNPJ_FUNDO", "DT_COMPTC", "VL_PATRIM_LIQ"})
    blc = pd.concat([le(z, n, COLS).assign(conf="CONFID" in n) for n in nomes if "_BLC_" in n or "CONFID" in n],
                    ignore_index=True)
    for df in (pl, blc):
        df["mes"] = df.DT_COMPTC.str[:7]
    links = blc.dropna(subset=["CNPJ_FUNDO_CLASSE_COTA"])[["mes", "CNPJ_FUNDO_CLASSE", "CNPJ_FUNDO_CLASSE_COTA"]]
    for m in meses:
        lk = links[links.mes == m]
        alvo = fecho(lk, peers, marca("cda", m) or ())
        b = blc[(blc.mes == m) & blc.CNPJ_FUNDO_CLASSE.isin(alvo)].copy()
        b["v"] = pd.to_numeric(b.VL_MERC_POS_FINAL, errors="coerce").fillna(0)
        b.loc[b.TP_APLIC.fillna("").str.contains(PASSIVO), "v"] *= -1

        def first(*cs):                           # primeira coluna preenchida entre as candidatas
            out = pd.Series(None, index=b.index, dtype=object)
            for c in cs:
                if c in b:
                    out = out.where(out.notna(), b[c])
            return out
        b["ativo"] = first("DS_ATIVO", "NM_FUNDO_CLASSE_SUBCLASSE_COTA", "DS_SWAP", "DS_ATIVO_EXTERIOR", "TP_TITPUB", "TP_ATIVO")
        b["emissor"] = first("EMISSOR", "CNPJ_EMISSOR", "CPF_CNPJ_EMISSOR")
        b["codigo"] = first("CD_ATIVO", "CD_ISIN", "CNPJ_FUNDO_CLASSE_COTA", "CD_SWAP", "CD_ATIVO_BV_MERC")
        b.loc[b.conf, "ativo"] = "(confidencial) " + b.loc[b.conf, "TP_APLIC"].fillna("")
        b["ativo"] = b.ativo.fillna(b.TP_ATIVO).fillna(b.TP_APLIC)
        venc = b.DT_VENC.fillna("").str[:7] if "DT_VENC" in b else ""
        b["ativo"] = (b.ativo.fillna("") + " " + b.emissor.fillna("").where(b.ativo.fillna("") != b.emissor.fillna(""), "")
                      + " " + venc).str.replace(r"\s+", " ", regex=True).str.strip()
        for c in NUM:
            b[c] = pd.to_numeric(b[c], errors="coerce") if c in b else None
        keep = ["CNPJ_FUNDO_CLASSE", "TP_APLIC", "TP_ATIVO", "ativo", "codigo", "emissor", "CNPJ_FUNDO_CLASSE_COTA", "v",
                "DT_VENC", "DS_INDEXADOR_POSFX", "GRAU_RISCO", "AG_RISCO", "PAIS", *NUM]
        b.reindex(columns=keep).to_parquet(CACHE / "rec" / f"cda_{m}.parquet", index=False)
        lk.drop(columns="mes").to_parquet(CACHE / "rec" / f"links_{m}.parquet", index=False)
        pl[pl.mes == m][["CNPJ_FUNDO_CLASSE", "VL_PATRIM_LIQ"]].to_parquet(CACHE / "rec" / f"pl_{m}.parquet", index=False)
        (CACHE / "rec" / f"cda_{m}.json").write_text(json.dumps({"v": VERSAO, "fundos": sorted(alvo)}))
    return url


def processa_diario(url, meses, peers):
    z = zipfile.ZipFile(baixa(url))
    cols = {"CNPJ_FUNDO_CLASSE", "CNPJ_FUNDO", "ID_SUBCLASSE", "DT_COMPTC", "VL_QUOTA", "VL_PATRIM_LIQ", "CAPTC_DIA",
            "RESG_DIA", "NR_COTST"}
    d = pd.concat([le(z, n, cols) for n in z.namelist()], ignore_index=True)
    alvo = set(peers).union(*[marca("diario", m) or set() for m in meses])   # mantem quem ja estava no recorte
    d = d[d.CNPJ_FUNDO_CLASSE.isin(alvo)]
    d = d[d.ID_SUBCLASSE.isna()] if "ID_SUBCLASSE" in d else d          # classe (nao subclasses)
    for m in meses:
        d[d.DT_COMPTC.str[:7] == m].to_parquet(CACHE / "rec" / f"diario_{m}.parquet", index=False)
        (CACHE / "rec" / f"diario_{m}.json").write_text(json.dumps({"v": VERSAO, "fundos": sorted(alvo)}))
    return url


def precisa(kind, m, peers):     # recorte em cache cobre os fundos pedidos?
    tem = marca(kind, m)
    if tem is None:
        return True
    if kind == "diario":
        return not set(peers) <= tem
    return not fecho(pd.read_parquet(CACHE / "rec" / f"links_{m}.parquet"), peers) <= tem


def _modificado(u):
    try:
        return u, pd.Timestamp(requests.head(u, timeout=60).headers["Last-Modified"]).timestamp()
    except (requests.RequestException, KeyError, ValueError):
        return u, None


def atualiza_publicacoes(cda, dia, log=print):
    """A CVM republica o arquivo do mes conforme os fundos entregam ou abrem a carteira (ha meses de prazo).
    Uma vez por dia confere a data de TODOS os arquivos mensais e rebaixa os que mudaram."""
    marca_dia = CACHE / f"verificado_{date.today()}.txt"
    if marca_dia.exists():
        return
    urls = [u for u in list(cda) + list(dia) if "HIST" not in u and (CACHE / u.rsplit("/", 1)[1]).exists()]
    with ThreadPoolExecutor(8) as ex:
        mudou = [u for u, t in ex.map(_modificado, urls) if t and t > (CACHE / u.rsplit("/", 1)[1]).stat().st_mtime]
    for u in mudou:
        (CACHE / u.rsplit("/", 1)[1]).unlink(missing_ok=True)
        for m in (cda.get(u) or dia.get(u)):
            for j in CACHE.glob(f"rec/*_{m}.json"):
                j.unlink(missing_ok=True)
    if mudou:
        log(f"{len(mudou)} arquivo(s) da CVM republicado(s) desde a ultima carga: baixando de novo")
    for velho in CACHE.glob("verificado_*.txt"):
        velho.unlink(missing_ok=True)
    marca_dia.write_text(f"{len(urls)} arquivos conferidos, {len(mudou)} atualizados")


def prepara(cnpjs, desde, workers=WORKERS, log=print):
    """Baixa/recorta da CVM so o que falta. Devolve (meses de carteira, {arquivo diario: meses})."""
    cda, dia = arquivos("CDA", desde), arquivos("INF_DIARIO", desde)
    atualiza_publicacoes(cda, dia, log)
    tarefas = ([(processa_cda, u, ms) for u, ms in cda.items() if any(precisa("cda", m, cnpjs) for m in ms)]
               + [(processa_diario, u, ms) for u, ms in dia.items() if any(precisa("diario", m, cnpjs) for m in ms)])

    def roda(i, fn):
        try:
            log(f"{i}/{len(tarefas)} {fn().rsplit('/', 1)[1]}")
        except Exception as e:                    # um arquivo ruim da CVM nao derruba o resto
            log(f"{i}/{len(tarefas)} ATENCAO: arquivo pulado ({type(e).__name__}: {e})")
    if workers > 1 and len(tarefas) > 1:
        with ProcessPoolExecutor(workers) as ex:
            fut = [ex.submit(f, u, ms, cnpjs) for f, u, ms in tarefas]
            for i, f in enumerate(fut, 1):
                roda(i, f.result)
    else:
        for i, (f, u, ms) in enumerate(tarefas, 1):
            roda(i, lambda: f(u, ms, cnpjs))
    meses = sorted(m for ms in cda.values() for m in ms if (CACHE / "rec" / f"cda_{m}.parquet").exists())
    return meses, dia


def fund_cat(name):
    n = (name or "").upper()
    return ("FIDC" if "FIDC" in n or "DIREITOS CREDIT" in n else "FIAGRO" if "FIAGRO" in n or "CADEIAS PRODUTIVAS" in n
            else "Cotas de Fundos")


def carrega_pos(meses):
    """Posicoes (so dos fundos recortados) + PL de todos, todos os meses, com a linha de ajuste carteira x PL."""
    pos = pd.concat([pd.read_parquet(CACHE / "rec" / f"cda_{m}.parquet").assign(mes=m) for m in meses], ignore_index=True)
    pl = pd.concat([pd.read_parquet(CACHE / "rec" / f"pl_{m}.parquet").assign(mes=m) for m in meses], ignore_index=True)
    pl = pl.drop_duplicates(["mes", "CNPJ_FUNDO_CLASSE"]).set_index(["mes", "CNPJ_FUNDO_CLASSE"]).VL_PATRIM_LIQ.astype(float)
    pos = pos.join(pl.rename("pl_fundo"), on=["mes", "CNPJ_FUNDO_CLASSE"])
    # a CDA nem sempre permite dar sinal certo a toda linha (ex.: compromissada tomada x doada): a diferenca
    # entre PL e soma da carteira vira uma linha explicita, em vez de inflar/encolher as categorias
    der = pos.TP_APLIC.fillna("").str.contains(DERIV)
    aj = (pos[~der].groupby(["mes", "CNPJ_FUNDO_CLASSE"]).agg(soma=("v", "sum"), pl_fundo=("pl_fundo", "first"))
          .assign(v=lambda a: a.pl_fundo - a.soma))
    aj = aj[(aj.v.abs() > 0.005 * aj.pl_fundo.abs())].reset_index().assign(
        TP_APLIC=AJUSTE, ativo="Diferença entre PL informado e soma da carteira (fonte CVM)")
    pos = pd.concat([pos, aj.drop(columns="soma")], ignore_index=True)
    return pos, pl


def lookthrough(nomes, meses, carga=None):
    """nomes: {nome: cnpj}. Uma linha por ativo final x caminho, com % do PL do fundo e movimentacao do mes."""
    pos, pl = carga or carrega_pos(meses)
    tem_cart = set(zip(pos.mes, pos.CNPJ_FUNDO_CLASSE))
    front = pd.DataFrame([dict(fundo=n, mes=m, veic=c, peso=1.0, caminho=n, cadeia=c)
                          for n, c in nomes.items() for m in meses if (m, c) in tem_cart])
    folhas = []
    for _ in range(12):
        if front.empty:
            break
        r = front.merge(pos, left_on=["mes", "veic"], right_on=["mes", "CNPJ_FUNDO_CLASSE"])
        f = r.peso / r.pl_fundo * 100              # R$ no veiculo -> % do PL do peer
        r["perc_pl"], r["compra_pct"], r["venda_pct"] = r.v * f, r.VL_AQUIS_NEGOC * f, r.VL_VENDA_NEGOC * f
        r["fator"] = f
        cota = r.CNPJ_FUNDO_CLASSE_COTA
        abre = cota.notna() & pd.Series([(m, c) in tem_cart and c not in cad for m, c, cad in zip(r.mes, cota, r.cadeia)],
                                        index=r.index, dtype=bool)
        nxt = r[abre]
        front = pd.DataFrame(dict(fundo=nxt.fundo, mes=nxt.mes, veic=nxt.CNPJ_FUNDO_CLASSE_COTA, peso=nxt.perc_pl / 100,
                                  caminho=nxt.caminho + " > " + nxt.ativo, cadeia=nxt.cadeia + "|" + nxt.CNPJ_FUNDO_CLASSE_COTA))
        folha = r[~abre].copy()
        eh_cota = folha.TP_APLIC.eq("Cotas de Fundos")
        folha["categoria"] = folha.TP_APLIC.where(~eh_cota, folha.ativo.map(fund_cat))
        folha["veiculo"] = folha.caminho.str.split(" > ").str[-1]
        folhas.append(folha)
    if not folhas:                              # CVM sem nenhuma carteira desse CNPJ no periodo
        return pd.DataFrame()
    df = pd.concat(folhas, ignore_index=True)
    df = df.rename(columns={"veic": "cnpj_veiculo", "TP_ATIVO": "tipo_ativo", "DT_VENC": "vencimento",
                            "DS_INDEXADOR_POSFX": "indexador", "PR_INDEXADOR_POSFX": "pct_indexador",
                            "PR_CUPOM_POSFX": "cupom", "PR_TAXA_PREFX": "taxa_pre", "GRAU_RISCO": "rating",
                            "AG_RISCO": "agencia", "PAIS": "pais", "v": "valor_veiculo", "QT_POS_FINAL": "quantidade",
                            "VL_CUSTO_POS_FINAL": "custo", "VL_AQUIS_NEGOC": "compras_veiculo",
                            "VL_VENDA_NEGOC": "vendas_veiculo", "pl_fundo": "pl_veiculo"})
    df = df[["fundo", "mes", "categoria", "ativo", "codigo", "emissor", "perc_pl", "compra_pct", "venda_pct", "veiculo",
             "cnpj_veiculo", "caminho", "tipo_ativo", "vencimento", "indexador", "pct_indexador", "cupom", "taxa_pre",
             "rating", "agencia", "pais", "quantidade", "valor_veiculo", "custo", "compras_veiculo", "vendas_veiculo",
             "pl_veiculo", "fator"]]
    df["categoria"] = df.categoria.fillna("Outros")
    df["codigo"] = df.codigo.str.strip()                  # CVM manda codigos com espacos (ex.: "ASER12  ")
    curto = df.ativo.fillna("").str.len() <= 4             # descricao generica ("DEB"): usa codigo + emissor
    df.loc[curto, "ativo"] = (df.codigo.fillna("") + " " + df.emissor.fillna("")).str.strip()[curto]
    df["derivativo"] = df.categoria.str.contains(DERIV)
    df["confidencial"] = df.ativo.str.startswith("(confidencial)", na=False)
    df["chave"] = df.cnpj_veiculo + "|" + df.codigo.fillna(df.ativo)
    return df


def mes_aberto(df, limite=5):
    """Por fundo, o mes mais recente em que a parte confidencial da carteira fica abaixo de `limite`% do PL."""
    t = df.groupby(["fundo", "mes"]).agg(conf=("perc_pl", lambda s: s[df.loc[s.index, "confidencial"]].sum()))
    ok = t[t.conf < limite].reset_index()
    return ok.groupby("fundo").mes.max()


def contribuicao(df):
    """Estimativa mensal de resultado por ativo pela variacao de preco das posicoes mantidas: quantidade do mes anterior
    x (preco unitario do mes - preco do mes anterior), em % do PL do peer. Preco unitario = valor / quantidade (CDA).
    Fica de fora: ganho de negociacao dentro do mes, carrego de caixa/compromissadas (sem quantidade) e ativos
    confidenciais. Cupons/amortizacoes pagos aparecem como queda de preco."""
    x = df[~df.derivativo & ~df.confidencial & ~df.categoria.eq(AJUSTE) & df.quantidade.gt(0)]
    k = x.groupby(["fundo", "mes", "chave"]).agg(v=("valor_veiculo", "sum"), qt=("quantidade", "sum"),
                                                  fator=("fator", "first"), categoria=("categoria", "first"),
                                                  ativo=("ativo", "first"), emissor=("emissor", "first")).reset_index()
    meses = sorted(df.mes.unique())
    seguinte = dict(zip(meses[:-1], meses[1:]))
    prev = k[["fundo", "mes", "chave", "v", "qt"]].assign(mes=k.mes.map(seguinte)).dropna(subset=["mes"])
    k = k.merge(prev.rename(columns={"v": "v_ant", "qt": "qt_ant"}), on=["fundo", "mes", "chave"])
    k["resultado_pct"] = k.qt_ant * (k.v / k.qt - k.v_ant / k.qt_ant) * k.fator
    return k


def _le_anbima(f):
    linhas = f.read_bytes().decode("latin1").splitlines()
    i = next(j for j, l in enumerate(linhas) if l.startswith("Código@"))
    a = pd.read_csv(io.StringIO("\n".join(linhas[i:])), sep="@", dtype=str)
    a.columns = ["codigo", "nome", "venc", "indice", "taxa_compra", "taxa_venda", "taxa_ind", "desvio", "int_min",
                 "int_max", "pu", "pct_par", "duration_du", "pct_reune", "ref_ntnb"][:a.shape[1]]
    for c in ["taxa_ind", "pu", "pct_par", "duration_du"]:
        a[c] = pd.to_numeric(a[c].str.replace(",", ".", regex=False), errors="coerce")
    ind = a.indice.fillna("").str.upper()
    a["tipo"] = pd.Series("Outros", index=a.index).mask(ind.str.contains("IGP"), "IGP-M +").mask(
        ind.str.contains("IPCA"), "IPCA +").mask(ind.str.contains("%") & ind.str.contains("DI"), "% do DI").mask(
        ind.str.match(r"^DI\s*\+"), "DI +")
    a["duration_anos"] = a.duration_du / 252
    return a.assign(data=pd.Timestamp(f"20{f.stem[-6:-4]}-{f.stem[-4:-2]}-{f.stem[-2:]}"))


def anbima_debentures():
    """Marcacao ANBIMA gratuita de debentures (taxa indicativa, PU, duration). O site so guarda ~15 dias uteis:
    cada execucao salva os arquivos da janela, entao o historico local cresce a partir de agora.
    Devolve (ultimo dia, historico de todos os dias guardados)."""
    hoje = pd.Timestamp.today().normalize()
    for dia in pd.bdate_range(hoje - pd.Timedelta(days=25), hoje):
        f = CACHE / f"anbima_db{dia:%y%m%d}.txt"
        if f.exists():
            continue
        try:
            r = requests.get(f"https://www.anbima.com.br/informacoes/merc-sec-debentures/arqs/db{dia:%y%m%d}.txt",
                             headers={"User-Agent": "Mozilla/5.0"}, timeout=60)
        except requests.RequestException:
            continue
        if r.status_code == 200 and "Código@".encode("latin1") in r.content:     # feriado/ainda nao publicado: pula
            f.write_bytes(r.content)
    fs = sorted(CACHE.glob("anbima_db*.txt"))
    if not fs:
        return pd.DataFrame(), pd.DataFrame()
    hist = pd.concat([_le_anbima(f) for f in fs], ignore_index=True)
    return hist[hist.data == hist.data.max()], hist


# ------------------------------------------------------------------ SND (debentures.com.br), FIDC (CVM) e Treasury: cobertura total
SND = "https://www.debentures.com.br/exploreosnd/consultaadados/emissoesdedebentures/"
UA = {"User-Agent": "Mozilla/5.0"}


def _get(url, tent=4, **kw):
    for t in range(tent):
        try:
            r = requests.get(url, headers=UA, timeout=300, **kw)
            r.raise_for_status()
            return r.content
        except requests.RequestException:
            if t == tent - 1:
                raise
            time.sleep(3 * (t + 1))


def _tabela_snd(conteudo, inicio):
    t = conteudo.decode("latin1").splitlines()
    i = next((j for j, l in enumerate(t) if l.startswith(inicio)), None)
    if i is None:
        return pd.DataFrame()
    a = pd.read_csv(io.StringIO("\n".join(t[i:])), sep="\t", dtype=str, quoting=3)
    a.columns = [c.strip() for c in a.columns]
    return a.apply(lambda s: s.str.strip())


def _diario(f):                    # arquivo que muda: rebaixa 1x por dia
    return not f.exists() or date.fromtimestamp(f.stat().st_mtime) < date.today()


def snd_caracteristicas():
    """Cadastro de TODAS as debentures do SND (ativas e ja vencidas/excluidas): indice, taxa de emissao, vencimento,
    periodicidade de juros, amortizacao e se e incentivada (Lei 12.431)."""
    partes = []
    for exc in ("False", "True"):
        f = CACHE / f"snd_caracteristicas_{exc}.txt"
        if _diario(f):
            try:
                f.write_bytes(_get(SND + f"caracteristicas_e.asp?tip_deb=publicas&op_exc={exc}&ativo="))
            except requests.RequestException:
                pass
        if f.exists():
            partes.append(_tabela_snd(f.read_bytes(), "Codigo do Ativo"))
    a = pd.concat(partes, ignore_index=True) if partes else pd.DataFrame()
    if not len(a):
        return pd.DataFrame()
    dt = lambda c: pd.to_datetime(a[c], format="%d/%m/%Y", errors="coerce")
    meses = lambda c, u: pd.to_numeric(a[c], errors="coerce") * np.where(a[u].eq("DIA"), 1 / 30, 1)
    out = pd.DataFrame({"codigo": a["Codigo do Ativo"], "empresa": a["Empresa"], "cnpj": a["CNPJ"],
                        "venc": dt("Data de Saida / Novo Vencimento").fillna(dt("Data de Vencimento")),
                        "indice": a["indice"].map(_sem_acento),
                        "pct": pd.to_numeric(a["Percentual Multiplicador/Rentabilidade"], errors="coerce"),
                        "juros": pd.to_numeric(a["Juros Criterio Novo - Taxa"].str.replace(",", "."), errors="coerce"),
                        "cada_juros": meses("Juros Criterio Novo - Cada", "Juros Criterio Novo - Unidade"),
                        "cada_amort": meses("Amortizacao - Cada", "Amortizacao - Unidade"),
                        "carencia": dt("Amortizacao - Carencia"),
                        "incentivada": a["Deb. Incent. (Lei 12.431)"].eq("S")})
    return out.drop_duplicates("codigo").set_index("codigo")


def snd_pu(datas, workers=4, log=print):
    """PU par (curva de emissao, com juros acumulados e amortizacoes) de todas as debentures em cada data.
    Um arquivo pequeno por data, guardado: so baixa datas novas."""
    vazio = lambda d: CACHE / f"snd_pu_{d:%Y%m%d}_vazio"        # feriado: marcado para nao pedir de novo

    def um(d):
        if vazio(d).exists():
            return None
        partes = []
        for e in ("False", "True"):             # ativas + ja vencidas/excluidas
            f = CACHE / f"snd_pu_{d:%Y%m%d}_{e}.txt"
            if not f.exists():
                try:
                    b = _get(SND + f"puhistorico_e.asp?op_exc={e}&ativo=&dt_ini={d:%d/%m/%Y}&dt_fim={d:%d/%m/%Y}")
                except requests.RequestException:
                    return None
                if e == "False" and b.count(b"\n") < 20:     # feriado / sem dados
                    if d < pd.Timestamp.today().normalize() - pd.Timedelta(days=7):
                        vazio(d).touch()
                    return None
                f.write_bytes(b)
            partes.append(_tabela_snd(f.read_bytes(), "Data do PU"))
        a = pd.concat(partes, ignore_index=True)
        a = a[a.iloc[:, 0].str.match(r"\d\d/\d\d/\d{4}", na=False)]
        num = lambda s: pd.to_numeric(s.str.replace(".", "", regex=False).str.replace(",", ".", regex=False), errors="coerce")
        return pd.DataFrame({"data": d, "codigo": a.iloc[:, 1].values, "pu_par": num(a.iloc[:, 5]).values})
    novas = [d for d in datas if not (CACHE / f"snd_pu_{d:%Y%m%d}_True.txt").exists() and not vazio(d).exists()]
    if novas:
        log(f"SND: PU par das debêntures em {len(novas)} fins de mês (só na 1a vez)...")
    with ThreadPoolExecutor(workers) as ex:
        out = [x for x in ex.map(um, datas) if x is not None]
    return pd.concat(out, ignore_index=True).drop_duplicates(["data", "codigo"]) if out else pd.DataFrame(columns=["data", "codigo", "pu_par"])


FIDC_URL = "https://dados.cvm.gov.br/dados/FIDC/DOC/INF_MENSAL/DADOS/"
_MEIO = [15, 45, 75, 105, 135, 165, 270, 540, 900, 1440]       # meio de cada faixa de prazo (dias) da tabela V


def _prazo_fidc(z, nome):
    t = pd.read_csv(z.open(nome), sep=";", encoding="latin1", dtype=str, quoting=3)
    c = "CNPJ_FUNDO_CLASSE" if "CNPJ_FUNDO_CLASSE" in t.columns else "CNPJ_FUNDO"
    b = t[[k for k in t.columns if re.match(r"TAB_V_A\d+_VL_PRAZO_VENC", k)]].apply(pd.to_numeric, errors="coerce").fillna(0).values
    tot = b.sum(1)
    return pd.DataFrame({"cnpj": t[c].values, "mes": t.DT_COMPTC.str[:7].values,
                         "anos": np.where(tot > 0, (b * _MEIO).sum(1) / np.where(tot > 0, tot, 1) / 365.25, np.nan)}).dropna()


def _series_fidc(z, n2, n3, n6):
    """Valor da cota (tab X_2), rentabilidade do mes (X_3) e desempenho esperado (X_6) de cada classe/serie de cada FIDC."""
    le = lambda n: pd.read_csv(z.open(n), sep=";", encoding="latin1", dtype=str, quoting=3)
    a, b = le(n2), le(n3)
    e = le(n6) if n6 in z.namelist() else pd.DataFrame(columns=list(a.columns[:4]) + ["TAB_X_CLASSE_SERIE", "TAB_X_PR_DESEMP_ESPERADO"])
    c = "CNPJ_FUNDO_CLASSE" if "CNPJ_FUNDO_CLASSE" in a.columns else "CNPJ_FUNDO"
    e = e.rename(columns={"CNPJ_FUNDO": c}) if c not in e.columns else e
    for t in (a, b, e):
        t["n"] = t.groupby([c, "DT_COMPTC", "TAB_X_CLASSE_SERIE"]).cumcount()
    k = [c, "DT_COMPTC", "TAB_X_CLASSE_SERIE", "n"]
    x = a.merge(b, on=k, how="outer").merge(e[k + ["TAB_X_PR_DESEMP_ESPERADO"]], on=k, how="left")
    num = lambda s: pd.to_numeric(s.str.replace(",", ".", regex=False), errors="coerce")
    return pd.DataFrame({"cnpj": x[c].values, "mes": x.DT_COMPTC.str[:7].values, "serie": x.TAB_X_CLASSE_SERIE.str.strip().values,
                         "cota": num(x.TAB_X_VL_COTA).values, "rent": num(x.TAB_X_VL_RENTAB_MES).values,
                         "esperado": num(x.TAB_X_PR_DESEMP_ESPERADO).values})


def fidc_mensal(desde, log=print):
    """Do informe mensal de FIDC da CVM: (1) prazo medio (anos) dos recebiveis a vencer de cada FIDC (tabela V) =
    duration estimada da cota; (2) valor da cota e rentabilidade de cada classe/serie (tabelas X_2/X_3) = spread da
    serie que o fundo tem. Guarda so o resultado (pequeno), nao os zips."""
    f, fs = CACHE / "fidc_prazo3.parquet", CACHE / "fidc_series3.parquet"
    tem = pd.read_parquet(f) if f.exists() else pd.DataFrame(columns=["cnpj", "mes", "anos"])
    ser = [pd.read_parquet(fs)] if fs.exists() else []
    feitos = set(tem.mes)
    hoje = pd.Timestamp.today()
    novos = []
    for a in range(int(desde[:4]), hoje.year + 1):
        ms = [f"{a}-{m:02d}" for m in range(1, 13) if pd.Timestamp(a, m, 1) <= hoje and f"{a}-{m:02d}" >= desde]
        falta = [m for m in ms if m not in feitos]
        if not falta:
            continue
        urls = [FIDC_URL + f"inf_mensal_fidc_{m.replace('-', '')}.zip" for m in falta] if a >= 2025 else [FIDC_URL + f"HIST/inf_mensal_fidc_{a}.zip"]
        for u in urls:
            try:
                z = zipfile.ZipFile(io.BytesIO(_get(u, tent=2)))
            except (requests.RequestException, zipfile.BadZipFile):
                continue
            log(f"FIDC: prazo e rentabilidade das séries ({u.rsplit('/', 1)[1]})...")
            nomes = z.namelist()
            for n in nomes:
                if "_tab_V_" in n:
                    novos.append(_prazo_fidc(z, n))
                if "_tab_X_2_" in n and n.replace("_X_2_", "_X_3_") in nomes:
                    ser.append(_series_fidc(z, n, n.replace("_X_2_", "_X_3_"), n.replace("_X_2_", "_X_6_")))
    if novos:
        tem = pd.concat([tem] + novos, ignore_index=True).drop_duplicates(["cnpj", "mes"], keep="last")
        tem.to_parquet(f, index=False)
        pd.concat(ser, ignore_index=True).drop_duplicates(["cnpj", "mes", "serie", "cota"], keep="last").to_parquet(fs, index=False)
    return tem, (pd.concat(ser, ignore_index=True) if ser else pd.DataFrame(columns=["cnpj", "mes", "serie", "cota", "rent", "esperado"]))


def treasury():
    """Curva de juros do Tesouro americano (treasury.gov, gratuito), no formato de tesouro() para usar em curva()."""
    tenor = {"1 Mo": 1 / 12, "3 Mo": .25, "6 Mo": .5, "1 Yr": 1, "2 Yr": 2, "3 Yr": 3, "5 Yr": 5, "7 Yr": 7, "10 Yr": 10, "20 Yr": 20, "30 Yr": 30}
    partes = []
    for a in range(int(DESDE[:4]), date.today().year + 1):
        f = CACHE / f"treasury_{a}.csv"
        if not f.exists() or (a == date.today().year and _diario(f)):
            try:
                f.write_bytes(_get("https://home.treasury.gov/resource-center/data-chart-center/interest-rates/daily-treasury-rates.csv/"
                                   f"{a}/all?type=daily_treasury_yield_curve&field_tdr_date_value={a}&page&_format=csv", tent=2))
            except requests.RequestException:
                continue
        if f.exists():
            partes.append(pd.read_csv(f))
    if not partes:
        return {}
    d = pd.concat(partes, ignore_index=True)
    d["data"] = pd.to_datetime(d.Date, format="%m/%d/%Y")
    d = d.drop_duplicates("data").sort_values("data")
    cols = [c for c in tenor if c in d.columns]
    x = np.array([tenor[c] for c in cols])
    por_dia = {r.data: (x[~np.isnan(v)], v[~np.isnan(v)]) for r, v in zip(d.itertuples(), d[cols].to_numpy(float)) if (~np.isnan(v)).sum() > 2}
    datas = sorted(por_dia)
    return {"UST": (np.array(datas, dtype="datetime64[ns]"), datas, por_dia)}


def sgs(codigo, desde):
    """Serie do Banco Central (SGS). Diarias: no maximo 10 anos por consulta."""
    out, hoje = [], pd.Timestamp.today()
    for a in range(int(desde[:4]), hoje.year + 1, 9):
        fim = min(pd.Timestamp(a + 8, 12, 31), hoje).strftime("%d/%m/%Y")
        for tent in range(5):
            try:
                out += requests.get(f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{codigo}/dados", timeout=60,
                                    params={"formato": "json", "dataInicial": f"01/01/{a}", "dataFinal": fim}).json()
                break
            except (requests.RequestException, ValueError):
                time.sleep(5 * (tent + 1))
    d = pd.DataFrame(out)
    return pd.Series(d.valor.astype(float).values, index=pd.to_datetime(d.data, dayfirst=True)).sort_index() if len(d) else pd.Series(dtype=float)


def cdi_mensal(desde, diario=False):
    s = sgs(12, desde) / 100
    return s if diario else (1 + s).groupby(s.index.strftime("%Y-%m")).prod() - 1


def ipca_indice():
    """Numero-indice do IPCA (IBGE/SIDRA), mensal. Usado no valor do contrato futuro DAP."""
    f = CACHE / "ipca_indice.json"
    if not f.exists() or f.stat().st_mtime < time.time() - 86400:
        try:
            r = requests.get("https://apisidra.ibge.gov.br/values/t/1737/n1/all/v/2266/p/all", timeout=120).json()
            f.write_text(json.dumps({x["D3C"]: x["V"] for x in r[1:]}))
        except (requests.RequestException, ValueError, KeyError):
            if not f.exists():
                return pd.Series(dtype=float)
    d = json.loads(f.read_text())
    s = pd.Series({f"{k[:4]}-{k[4:]}": float(v) for k, v in d.items() if v not in ("...", "-", "")})
    return s.sort_index()


def tesouro():
    """Curvas do Tesouro Direto (gratuito, diario desde 2004): prefixado (LTN) e IPCA+ (NTN-B principal), taxa x prazo."""
    f = CACHE / "tesouro_direto.csv"
    if not f.exists() or f.stat().st_mtime < time.time() - 86400:
        try:
            r = requests.get("https://www.tesourotransparente.gov.br/ckan/dataset/df56aa42-484a-4a59-8184-7676580c81e3/"
                             "resource/796d2059-14e9-44e3-80c9-2d9e30b405c1/download/PrecoTaxaTesouroDireto.csv", timeout=300)
            r.raise_for_status()
            f.write_bytes(r.content)
        except requests.RequestException:
            if not f.exists():
                return {}
    d = pd.read_csv(f, sep=";", decimal=",")
    d = d[d["Tipo Titulo"].isin(["Tesouro Prefixado", "Tesouro IPCA+"])].copy()
    d["data"] = pd.to_datetime(d["Data Base"], dayfirst=True)
    d["anos"] = (pd.to_datetime(d["Data Vencimento"], dayfirst=True) - d.data).dt.days / 365.25
    d["taxa"] = d[["Taxa Compra Manha", "Taxa Venda Manha"]].replace(0, np.nan).mean(axis=1)
    d["tipo"] = np.where(d["Tipo Titulo"].eq("Tesouro Prefixado"), "PRE", "IPCA")
    d = d.dropna(subset=["taxa"]).sort_values(["tipo", "data", "anos"])
    out = {}
    for t, g in d.groupby("tipo"):
        por_dia = {k: (x.anos.values, x.taxa.values) for k, x in g.groupby("data")}
        datas = sorted(por_dia)
        out[t] = (np.array(datas, dtype="datetime64[ns]"), datas, por_dia)
    return out


def curva(curvas, tipo, dia, anos):
    """Taxa (% a.a.) da curva `tipo` ("PRE" ou "IPCA") no ultimo dia util <= `dia`, interpolada no prazo `anos`."""
    if tipo not in curvas:
        return np.full(np.shape(anos), np.nan)
    arr, datas, por_dia = curvas[tipo]
    i = np.searchsorted(arr, np.datetime64(pd.Timestamp(dia)), side="right") - 1
    if i < 0:
        return np.full(np.shape(anos), np.nan)
    x, y = por_dia[datas[i]]
    return np.interp(np.clip(anos, x.min(), x.max()), x, y)


def _sem_acento(t):
    import unicodedata
    return unicodedata.normalize("NFKD", str(t)).encode("ascii", "ignore").decode().upper()


_GRUPOS_RE = [(re.compile(k), v) for k, v in GRUPOS.items()]
_SUFIXOS = re.compile(r"(?<![\w.])(S\.A\.?|S/A|SA|LTDA|SPE|PARTICIPACOES|HOLDING|COMPANHIA|CIA|EMPREENDIMENTOS|E|DE|DO|DA|DOS|DAS|EM)(?![\w/])")


def grupo(emissor, categoria=""):
    """Grupo economico pelo nome do emissor (mapa GRUPOS); sem mapa, o proprio emissor sem sufixos."""
    if categoria in ("Operações Compromissadas",):
        return "Compromissadas"
    if categoria == "Títulos Públicos":
        return "Tesouro Nacional"
    if categoria in ("Disponibilidades", "Valores a receber", "Valores a pagar", AJUSTE):
        return "Caixa e outros"
    if not isinstance(emissor, str) or not emissor.strip():
        return None
    n = _sem_acento(emissor)
    for rx, g in _GRUPOS_RE:
        if rx.search(n):
            return g
    base = " ".join(w for w in _SUFIXOS.sub(" ", n).split() if not w.isdigit())
    return base.title()[:40] or emissor


def diario(nomes, dia):
    nome = {c: n for n, c in nomes.items()}
    d = pd.concat([pd.read_parquet(CACHE / "rec" / f"diario_{m}.parquet") for ms in dia.values() for m in ms
                   if (CACHE / "rec" / f"diario_{m}.parquet").exists()], ignore_index=True)
    d = d[d.CNPJ_FUNDO_CLASSE.isin(nome)]
    for c in ["VL_QUOTA", "VL_PATRIM_LIQ", "CAPTC_DIA", "RESG_DIA", "NR_COTST"]:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    d = d.drop_duplicates(["CNPJ_FUNDO_CLASSE", "DT_COMPTC"], keep="last").sort_values("DT_COMPTC")
    return d.assign(data=pd.to_datetime(d.DT_COMPTC), mes=d.DT_COMPTC.str[:7], fundo=d.CNPJ_FUNDO_CLASSE.map(nome))


def rentabilidade(d, cdi):
    g = d.groupby(["fundo", "mes"])
    rent = pd.DataFrame({"cota_fim": g.VL_QUOTA.last(), "pl_fim": g.VL_PATRIM_LIQ.last(),
                         "captacao_liquida": g.CAPTC_DIA.sum() - g.RESG_DIA.sum(), "cotistas": g.NR_COTST.last()})
    rent["retorno"] = rent.groupby(level=0).cota_fim.pct_change()
    rent["cdi"] = rent.index.get_level_values(1).map(cdi)
    rent["pct_cdi"] = rent.retorno / rent.cdi
    return rent.reset_index()


def excel(df, rent, ordem, out):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.chart import AreaChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils.indexed_list import IndexedList
    limpa = lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v
    df = df.copy()
    for c in df.select_dtypes("object"):
        df[c] = df[c].map(limpa)
    B, N = Font(name="Arial", size=10, bold=True), Font(name="Arial", size=10)
    HF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    FILL, MFILL = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="D9E1F2")
    PCT, PCT2 = '0.0000"%";(0.0000"%")', '0.00"%";(0.00"%")'

    def put(ws, vals, font=None, fill=None, fmt=None):
        ws.append(vals)
        r = ws._current_row                       # O(1); ws.max_row recontaria a planilha toda
        for j in range(1, len(vals) + 1):
            c = ws.cell(r, j)
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if fmt and j in fmt:
                c.number_format = fmt[j]
        return r

    def header(ws, cols, widths):
        put(ws, cols, HF, FILL)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w

    wb = Workbook()
    wb._fonts = IndexedList([N])                  # Arial 10 como fonte padrao
    wb.remove(wb.active)
    for nm, g in df.groupby("fundo", sort=False):
        ws = wb.create_sheet(nm[:31])
        header(ws, ["Categoria / Ativo", "Codigo", "% PL " + nm, "% da categoria", "Veiculo", "Emissor"], (70, 18, 14, 14, 50, 40))
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.freeze_panes = "A2"
        for ym, gm in sorted(g.groupby("mes"), reverse=True):
            put(ws, [ym, None, gm.perc_pl.sum()], B, MFILL, {3: PCT})
            ats = gm.groupby(["categoria", "ativo", "codigo"], dropna=False).agg(
                perc_pl=("perc_pl", "sum"), veiculo=("veiculo", "first"), emissor=("emissor", "first"))
            for cat, tot in gm.groupby("categoria").perc_pl.sum().sort_values(ascending=False).items():
                r = put(ws, ["   " + cat, None, tot], B, None, {3: PCT})
                ws.row_dimensions[r].outline_level = 1
                for (_, ativo, cod), a in ats.loc[[cat]].sort_values("perc_pl", ascending=False).iterrows():
                    last = put(ws, ["      " + ativo, None if pd.isna(cod) else cod, a.perc_pl,
                                    a.perc_pl / tot * 100 if tot else None, a.veiculo, a.emissor], fmt={3: PCT, 4: PCT})
                ws.row_dimensions.group(r + 1, last, outline_level=2, hidden=True)

    CURTO = {"Depósitos a prazo e outros títulos de IF": "Depósitos a prazo / IF",
             "Outros valores mobiliários registrados na CVM objeto de oferta pública": "Outros VM (CVM)",
             "Títulos ligados ao agronegócio": "Títulos agro", "Títulos de Crédito Privado": "Crédito privado (CCB/NC)"}
    CORES, OUTROS_COR = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7"], "898781"
    cat_ = df.categoria.map(lambda c: CURTO.get(c, c))
    CONTABIL = r"(?i)pagar|receber|obriga|termo|disponibilidade|exigibilidade|swap|confidencial|ajuste"
    TOP = (df.assign(c=cat_)[~cat_.str.contains(CONTABIL) & ~df.derivativo].groupby(["fundo", "mes", "c"]).perc_pl.sum()
           .groupby("c").mean().sort_values(ascending=False).index[:7].tolist())
    SERIES = TOP + ["Outros"]
    ev = df[~df.derivativo].assign(c=cat_.where(cat_.isin(TOP), "Outros")).pivot_table(
        index=["fundo", "mes"], columns="c", values="perc_pl", aggfunc="sum", fill_value=0).reindex(columns=SERIES, fill_value=0)
    ws, gr = wb.create_sheet("evolucao", 0), wb.create_sheet("graficos", 0)
    put(gr, ["Composicao por categoria (% do PL, look-through, dados CVM). Numeros em 'evolucao'."], B)
    for k, nm in enumerate(n for n in ordem if n in ev.index.get_level_values(0)):
        row0 = put(ws, [nm], B)
        put(ws, ["mes"] + SERIES, HF, FILL)
        for ym, r in ev.loc[nm].iterrows():
            last = put(ws, [ym] + [float(v) for v in r], fmt={j: PCT2 for j in range(2, len(SERIES) + 2)})
        ws.append([])
        ws.append([])
        ch = AreaChart()
        ch.grouping, ch.title, ch.y_axis.title = "stacked", f"{nm} — % do PL por categoria", "% do PL"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorGridlines = 0, 100, None
        ch.x_axis.tickLblSkip, ch.x_axis.delete, ch.y_axis.delete = 12, False, False
        ch.add_data(Reference(ws, min_col=2, max_col=1 + len(SERIES), min_row=row0 + 1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=row0 + 2, max_row=last))
        for s, cor in zip(ch.series, CORES + [OUTROS_COR]):
            s.graphicalProperties.solidFill, s.graphicalProperties.line.solidFill = cor, "FCFCFB"
        ch.legend.position, ch.width, ch.height = "r", 26, 9
        gr.add_chart(ch, f"A{3 + k * 19}")
    ws.column_dimensions["A"].width = 12

    ws = wb.create_sheet("rentabilidade", 2)
    header(ws, ["fundo", "mes", "retorno", "CDI", "% CDI", "cota fim", "PL fim", "captacao liquida", "cotistas"],
           (16, 9, 11, 11, 11, 14, 18, 18, 10))
    for r in rent.itertuples(index=False):
        put(ws, [r.fundo, r.mes, r.retorno, r.cdi, r.pct_cdi, r.cota_fim, r.pl_fim, r.captacao_liquida, r.cotistas],
            fmt={3: "0.00%", 4: "0.00%", 5: "0.0%", 7: "#,##0", 8: "#,##0;(#,##0)"})

    ult = df[df.mes == df.fundo.map(df.groupby("fundo").mes.max())]
    top10 = (ult.groupby(["categoria", "fundo", "mes", "ativo", "codigo"], dropna=False).perc_pl.sum().reset_index()
             .sort_values(["categoria", "perc_pl"], ascending=[True, False]).groupby("categoria").head(10))
    top10.insert(1, "rank", top10.groupby("categoria").cumcount() + 1)
    ws = wb.create_sheet("top10", 3)
    header(ws, ["categoria", "rank", "fundo", "mes", "ativo", "codigo", "% PL do fundo"], (34, 6, 16, 9, 70, 18, 14))
    for r in top10.itertuples(index=False):
        put(ws, [None if pd.isna(x) else x for x in r], fmt={7: PCT})
    wb.save(out)
    df.drop(columns=["chave", "fator"]).to_csv(out.replace(".xlsx", "_dados.csv"), index=False, sep=";", decimal=",",
                                               encoding="utf-8-sig")


lt = sys.modules[__name__]          # o servidor chama as funcoes de dados como lt.funcao

# ====================================================================== servidor + pagina

import json
import math
import os
import re
import threading
from datetime import date

import numpy as np
import pandas as pd
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response


app = FastAPI(title="Painel de fundos")

CAT_CORES = ["#2a78d6", "#1baf7a", "#eda100", "#e87ba4", "#4a3aa7", "#008300", "#e34948", "#7ba7d9"]
NOSSAS_CORES = ["#eb6834", "#c24f1d", "#f29a6b", "#a8401a", "#f5b38d"]
GENERICO = (r"(?i)\b(FUNDO DE INVESTIMENTO|FUNDO|EM DIREITOS CREDIT[OÓ]RIOS|DE INVESTIMENTO|FIDC|FIC|FI|COTAS|MULTIMERCADO|"
            r"RENDA FIXA|CR[EÉ]DITO PRIVADO|RESPONSABILIDADE LIMITADA|RESP LTDA|N[AÃ]O PADRONIZADO|NP|LONGO PRAZO|LP)\b")
METODO = 13                                     # suba quando mudar o calculo -> refaz o cache do dia
CAIXA = ("Caixa",)
CAIXA_CVM = ("Disponibilidades", "Operações Compromissadas", "Valores a receber", "Valores a pagar", lt.AJUSTE)
TIPO_CASA = {"CDB/ RDB": "CDB", "CDB Vinculado": "CDB", "DPGE": "DPGE", "FI Imobiliário": "FII", "FI Participações": "FIP",
             "Fundos de Índice": "ETF", "Fundos Offshore": "Fundos offshore", "CCB": "CCB", "CCCB": "CCB", "LCA": "LCA", "CPR": "CPR",
             "CDCA": "CDCA", "CCI": "CCI", "NCA": "NCA", "Letra de Câmbio/ Letra Hipotecária/ Letra Imobiliária": "LCI/LH"}
CURTO_CVM = {"Depósitos a prazo e outros títulos de IF": "Outros títulos de IF", "Títulos ligados ao agronegócio": "Agro (outros)",
             "Outros valores mobiliários registrados na CVM objeto de oferta pública": "Outros valores mobiliários",
             "Investimento no Exterior": "Exterior (outros)", "Títulos de Crédito Privado": "Crédito privado (outros)",
             "Cotas de Fundos": "Cotas de fundos"}
SEM_SPREAD = ("Ações", "FII", "FIP", "ETF", "FIAGRO", "Confidencial")
MERCADO = ("ANBIMA", "SND + preço do fundo")            # renda variavel: spread e duration nao se aplicam
FUNDOS = ("FIDC", "FII", "FIP", "ETF", "FIAGRO", "Cotas de fundos", "Fundos offshore")
USD = ("Bonds", "Fundos offshore", "Exterior (outros)")
TP = [("LFT", r"BRSTNCLF|FINANCEIRAS DO TESOURO|LFT"), ("NTN-B", r"BRSTNCNTB|SERIE B|NTN-B"),
      ("NTN-F", r"BRSTNCNTF|SERIE F|NTN-F"), ("LTN", r"BRSTNCLTN|LETRAS DO TESOURO NACIONAL|LTN")]
MESES_COD = "FGHJKMNQUVXZ"
PRE = {"status": "parado", "log": [], "erro": None, "pasta": None}


# ------------------------------------------------------------------ utilidades
def limpo(o):
    if isinstance(o, dict):
        return {str(k): limpo(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [limpo(v) for v in o]
    if isinstance(o, np.integer):
        return int(o)
    if isinstance(o, (np.floating, float)):
        return None if (math.isnan(o) or math.isinf(o)) else round(float(o), 6)
    if isinstance(o, pd.Timestamp):
        return None if pd.isna(o) else o.strftime("%Y-%m-%d")
    if o is None or isinstance(o, (str, bool)):
        return o
    try:
        return None if pd.isna(o) else o
    except (TypeError, ValueError):
        return o


def wavg(v, w):
    v, w = np.asarray(v, float), np.asarray(w, float)
    ok = ~np.isnan(v) & (w > 0)
    return float(np.average(v[ok], weights=w[ok])) if ok.any() else None


def fim_mes(m):
    return pd.Timestamp(m) + pd.offsets.MonthEnd(0)


def contrato(texto):
    """Tipo, vencimento (ano, mes) de um futuro a partir do codigo/descricao da CDA (ex.: DI1FUTF29, FUT DAP/Q30)."""
    t = lt._sem_acento(texto)
    tipo = ("DAP" if re.search(r"DAP|CUPOM DE IPCA|DI X IPCA", t) else "DI1" if re.search(r"DI1|DI DE 1 DIA", t)
            else "WDO" if "WDO" in t else "DOL" if re.search(r"DOL|DOLAR", t) else "DDI" if re.search(r"DDI|FRC|CUPOM CAMBIAL", t)
            else "IND" if re.search(r"\bIND|WIN|IBOV", t) else "OUTRO")
    m = re.search(r"(?:DI1|DAP|DOL|WDO|DDI|FRC|IND|WIN)\s*(?:FUT)?\s*/?\s*([FGHJKMNQUVXZ])(\d{2})\b", t)
    return tipo, ((2000 + int(m.group(2)), MESES_COD.index(m.group(1)) + 1) if m else None)


# ------------------------------------------------------------------ fluxo de caixa: duration e taxa implicita pelo preco
def _grade(T, p_jur, p_amo, t_car):
    """Fluxos por unidade de principal numa grade mensal contada a partir do vencimento (vetorizado).
    T: anos ate o vencimento; juros a cada p_jur anos; amortizacao constante a cada p_amo anos (0 = no vencimento)
    a partir de t_car anos. Devolve (prazos, peso_juros, principal, periodo_juros)."""
    T = np.clip(np.nan_to_num(np.atleast_1d(np.asarray(T, float)), nan=1.0), 1 / 12, 40)
    col = lambda v, pad: np.broadcast_to(np.nan_to_num(np.asarray(v, float), nan=pad), T.shape)[:, None]
    j = np.arange(int(np.ceil(T.max() * 12)) + 1)[None, :]
    tau = T[:, None] - j / 12
    ok = tau > 1e-6
    kj = np.clip(np.round(col(p_jur, 0.5) * 12), 1, 600)
    ka = np.round(col(p_amo, 0) * 12)
    am = np.where(ka > 0, (j % np.maximum(ka, 1) == 0) & (tau >= col(t_car, 0) - 1e-6), j == 0) & ok
    am |= (j == 0) & ~am.any(1, keepdims=True)
    princ = am / am.sum(1, keepdims=True)
    saldo = np.cumsum(princ, 1)                       # saldo devedor antes do pagamento em tau
    return np.where(ok, tau, 1.0), ((j % kj == 0) & ok) * saldo, princ, kj / 12


def pv(tau, wj, princ, per, y_cup, y):
    """Valor presente por unidade de principal (e soma de prazo x VP): juros a taxa do titulo y_cup, desconto a y (% a.a.)."""
    fl = wj * ((1 + np.asarray(y_cup, float)[:, None] / 100) ** per - 1) + princ
    df = (1 + np.asarray(y, float)[:, None] / 100) ** -tau
    return (fl * df).sum(1), (fl * df * tau).sum(1)


def taxa_e_duration(T, p_jur, p_amo, t_car, base, s0, pct_par):
    """Taxa de mercado s (mesma unidade de s0) que leva o preco a pct_par do PU par, e a duration (anos) nessa taxa.
    y = (1+base)(1+s)-1: base = CDI para DI+, 0 para IPCA+/pre (s ja e a taxa). pct_par NaN -> fica s0."""
    tau, wj, princ, per = _grade(T, p_jur, p_amo, t_car)
    n = len(tau)
    base, s0 = np.broadcast_to(np.asarray(base, float), (n,)), np.broadcast_to(np.asarray(s0, float), (n,)).copy()
    alvo = np.broadcast_to(np.asarray(pct_par, float), (n,))
    y = lambda s, b: ((1 + b / 100) * (1 + s / 100) - 1) * 100
    s = s0.copy()
    r = np.where(np.isfinite(alvo) & np.isfinite(s0))[0]
    if len(r):                                       # bissecao: o preco cai quando a taxa sobe
        a = (tau[r], wj[r], princ[r], per[r])
        v0, _ = pv(*a, y(s0[r], base[r]), y(s0[r], base[r]))
        lo, hi = np.full(len(r), -10.0), np.full(len(r), 80.0)
        for _ in range(30):
            mid = (lo + hi) / 2
            v, _ = pv(*a, y(s0[r], base[r]), y(mid, base[r]))
            caro = v / v0 > alvo[r]
            lo, hi = np.where(caro, mid, lo), np.where(caro, hi, mid)
        s[r] = (lo + hi) / 2
    v, vt = pv(tau, wj, princ, per, y(np.nan_to_num(s0), base), y(np.nan_to_num(s), base))
    return s, vt / v


# ------------------------------------------------------------------ calculo por fundo
class Contexto:
    """Dados de mercado compartilhados por todos os fundos."""

    def __init__(self, log, meses):
        log("Baixando curvas (Tesouro Direto), IPCA (IBGE), PTAX e CDI (Banco Central) e marcação ANBIMA...")
        self.curvas = lt.tesouro()
        self.ipca = lt.ipca_indice()
        self.ptax = lt.sgs(1, lt.DESDE)
        self.cdi_d = lt.cdi_mensal(lt.DESDE, diario=True)
        self.cdi_m = (1 + self.cdi_d).groupby(self.cdi_d.index.strftime("%Y-%m")).prod() - 1
        self.an_u, self.an_h = lt.anbima_debentures()
        self.an_dias = sorted(self.an_h.data.unique()) if len(self.an_h) else []
        log("Cadastro e PU par das debêntures (SND), prazo dos FIDCs (CVM) e Treasury (EUA)...")
        self.snd = lt.snd_caracteristicas()
        hoje = pd.Timestamp.today().normalize() - pd.offsets.BDay(1)
        alvo = {m: min(pd.Timestamp(m) + pd.offsets.BMonthEnd(0), hoje) for m in meses}
        pu = lt.snd_pu(sorted(set(alvo.values())), log=log)
        falta = {m: d - pd.offsets.BDay(1) for m, d in alvo.items() if d not in set(pu.data)}      # feriado no fim do mes
        if falta:
            pu = pd.concat([pu, lt.snd_pu(sorted(set(falta.values())), log=log)], ignore_index=True)
            alvo.update(falta)
        self.pu_dia = alvo
        self.pu = {d: x.drop_duplicates("codigo").set_index("codigo").pu_par for d, x in pu.groupby("data")}
        f, ser = lt.fidc_mensal(lt.DESDE, log)
        self.fidc = {c: (x.mes.values, x.anos.values) for c, x in f.sort_values("mes").groupby("cnpj")}
        ser = ser[(ser.cota > 0) & ser.rent.notna()].sort_values(["cnpj", "serie", "mes"])
        e = (1 + ser.rent / 100) / (1 + ser.mes.map(self.cdi_m)) - 1
        primeiro = ser.groupby(["cnpj", "serie"]).cumcount() == 0             # 1o mes da serie e parcial
        ser = ser.assign(e=e.where(~primeiro & (e > -0.03) & (e < 0.2)))       # < -3% no mes: amortizacao/evento pontual
        ser["spread"] = ((1 + ser.groupby(["cnpj", "serie"]).e.transform(lambda x: x.rolling(6, min_periods=2).median())) ** 12 - 1) * 100
        esp = ((1 + ser.esperado / 100) / (1 + ser.mes.map(self.cdi_m))) ** 12 * 100 - 100        # benchmark da serie (X_6)
        ruim = ~((ser.spread > -30) & (ser.spread < 40))                 # perda recorrente aparece; absurdo vira benchmark
        ser["spread"] = ser.spread.mask(ruim, esp.where((ser.esperado > 0) & (esp > -5) & (esp < 30)))
        ser = ser.dropna(subset=["spread"])
        sr = ser.serie.str.contains("Senior|Sênior", case=False, na=False)
        self.fidc_ser = {k: (x.cota.values, x.spread.values, float(x.spread[sr[x.index]].median()) if sr[x.index].any() else float(x.spread.median()))
                         for k, x in ser.groupby(["cnpj", "mes"])}
        self.ust = lt.treasury()
        px = self.ptax.groupby(self.ptax.index.strftime("%Y-%m")).last()
        self.fx_m = px / px.shift() - 1
        u = pd.Series({m: float(lt.curva(self.ust, "UST", fim_mes(m), np.array([1.0]))[0]) / 1200 for m in px.index}) if self.ust else pd.Series(dtype=float)
        self.ust_m = u.shift()
        self._taxas = {}
        self.emissores = {}                       # (mes, grupo) -> soma spread*peso, duration*peso, peso

    def pu_par(self, m):
        return self.pu.get(self.pu_dia.get(m), pd.Series(dtype=float))

    def proj_ipca(self, m):
        """IPCA mensal projetado (media dos 12 meses anteriores): o PU par do SND so usa o indice ja divulgado."""
        x = self.ipca[self.ipca.index < m]
        return float((x.iloc[-1] / x.iloc[-13]) ** (1 / 12) - 1) if len(x) > 13 else 0.004

    def fidc_anos(self, cnpj, m):
        x = self.fidc.get(cnpj)
        if x is None:
            return np.nan
        i = np.searchsorted(x[0], m, side="right") - 1
        return float(x[1][max(i, 0)])

    def fidc_spread(self, cnpj, m, pu):
        """Spread da serie do FIDC que o fundo tem: acha a serie pelo valor da cota (tab X_2) e usa a mediana de 6 meses
        da rentabilidade acima do CDI (tab X_3; benchmark X_6 se ela for anomala). Sem serie com a mesma cota: seniores."""
        for k in range(4):                                  # informe do mes ou de ate 3 meses antes
            x = self.fidc_ser.get((cnpj, (pd.Timestamp(m) - pd.DateOffset(months=k)).strftime("%Y-%m")))
            if x is not None:
                cota, spr, senior = x
                if np.isfinite(pu) and pu > 0:
                    dif = np.abs(cota / pu - 1)
                    if dif.min() < 0.02:
                        return float(spr[dif.argmin()])
                return senior
        return np.nan

    def taxa(self, m, cods, T, pj, pa, car, base, s0, pp_):
        """taxa_e_duration com memoria: o mesmo ativo no mesmo mes e preco aparece em varios fundos."""
        ch = [(m, c, round(float(p), 5) if np.isfinite(p) else None, round(float(b), 4)) for c, p, b in zip(cods, pp_, s0)]
        novos = {k: i for i, k in enumerate(ch) if k not in self._taxas}
        if novos:
            i = np.array(list(novos.values()))
            s, d = taxa_e_duration(T[i], pj[i], pa[i], car[i], base[i], s0[i], pp_[i])
            self._taxas.update(zip(novos, zip(s, d)))
        r = np.array([self._taxas[k] for k in ch], dtype=float).reshape(-1, 2)
        return r[:, 0], r[:, 1]

    def cdi_aa(self, m):
        x = self.cdi_d[self.cdi_d.index <= fim_mes(m)]
        return float(((1 + x.iloc[-1]) ** 252 - 1) * 100) if len(x) else np.nan

    def ptax_em(self, m):
        x = self.ptax[self.ptax.index <= fim_mes(m)]
        return float(x.iloc[-1]) if len(x) else np.nan

    def anbima_mes(self, m, recentes):
        """Marcacao ANBIMA para a carteira do mes: o dia guardado mais proximo do fim do mes (ate 10 dias);
        nos meses mais recentes, o ultimo dia disponivel."""
        if not self.an_dias:
            return None, None
        alvo = fim_mes(m)
        perto = [d for d in self.an_dias if abs((pd.Timestamp(d) - alvo).days) <= 10]
        d = min(perto, key=lambda d: abs((pd.Timestamp(d) - alvo).days)) if perto else (self.an_dias[-1] if m in recentes else None)
        return (d, self.an_h[self.an_h.data == d].drop_duplicates("codigo").set_index("codigo")) if d is not None else (None, None)


def tipo_tp(g):
    s = (g.codigo.fillna("") + " " + g.ativo.fillna("")).map(lt._sem_acento)
    out = pd.Series(None, index=g.index, dtype=object)
    for nome, rx in TP:
        out = out.mask(out.isna() & s.str.contains(rx), nome)
    return out


def vencimento(g):
    """Vencimento informado ou lido do nome do ativo (CRI: 'CRI:APCS:150629', '... / 25/09/2029 / ...', '... 2029-05')."""
    a = g.ativo.fillna("")
    v = pd.to_datetime(g.vencimento, errors="coerce")
    for rx, fmt in [(r"(\d{2}/\d{2}/\d{4})", "%d/%m/%Y"), (r"CR[AI]:[A-Z0-9]+:(\d{6})\b", "%d%m%y"), (r"\b((?:19|20)\d{2}-\d{2})$", "%Y-%m")]:
        v = v.fillna(pd.to_datetime(a.str.extract(rx)[0], format=fmt, errors="coerce"))
    return v


def casa(df, ctx):
    """Categorias da casa: Caixa, LF, LFSN, LFSC, DEB, DEBIN, Bonds, CRA, CRI, NC, FIDC. O que nao se encaixa fica com o
    nome do tipo/categoria da CVM (CDB, Títulos Públicos, FII, ...)."""
    t, c = df.tipo_ativo.fillna(""), df.categoria
    a = df.ativo.fillna("").map(lt._sem_acento)
    inc = df.codigo.map(ctx.snd.incentivada) if len(ctx.snd) else pd.Series(np.nan, index=df.index)
    out = t.map(TIPO_CASA).fillna(c.map(CURTO_CVM)).fillna(c)
    fundo = t.isin(["FIDC", "FI Imobiliário", "Fundos de Índice", "FI Participações", "Fundos Offshore"]) | c.isin(["FIDC", "Cotas de Fundos", "FIAGRO"])
    out = out.mask(c.eq("FIAGRO"), "FIAGRO").mask(t.str.startswith("Aç") | t.str.contains("subscri"), "Ações")
    out = out.mask(t.isin(["Bonds e Treasury", "Título da Dívida Externa"]), "Bonds")
    out = out.mask(t.eq("FIDC") | c.eq("FIDC"), "FIDC")
    out = out.mask(~fundo & (t.str.contains("Nota Promiss|Nota Comercial") | a.str.match(r"NOTA COMERCIAL")), "NC")
    out = out.mask(~fundo & (t.str.contains("imobili.rios") | a.str.contains(r"\bCRI\b")), "CRI")
    out = out.mask(~fundo & (t.eq("CRA") | a.str.match(r"CRA\b")), "CRA")
    deb = ~fundo & (t.str.startswith("Debênture") | a.str.match(r"DEBENTURE") | inc.notna())
    out = out.mask(deb, np.where(inc.eq(True), "DEBIN", "DEB"))
    lf = t.eq("Letra Financeira")
    venc = pd.to_datetime(df.vencimento, errors="coerce")
    chave = df.emissor.fillna("") + "|" + df.vencimento.astype(str)
    primeiro = pd.to_datetime(df.mes.where(lf)).groupby(chave).transform("min")
    longo = (venc - primeiro).dt.days / 365.25 > 6          # LF subordinada nivel II: prazo longo (> 6 anos), nao perpetua
    out = out.mask(lf, np.where(venc.dt.year >= 2049, "LFSC", np.where(longo, "LFSN", "LF")))
    out = out.mask(c.eq("Debêntures") & ~out.isin(["DEB", "DEBIN"]), "DEB")
    contabil = c.isin(CAIXA_CVM) | (c.str.contains(lt.PASSIVO, na=False) & ~df.derivativo)
    out = out.mask(contabil, "Caixa").mask(c.eq("Títulos Públicos"), "Títulos Públicos")
    return out.mask(df.confidencial & ~df.derivativo, "Confidencial")       # CVM so informa o tipo, sem o ativo


def rotulo(g, venc, tp):
    """Nome curto 'codigo - mm/aaaa': debenture/acao pelo ticker, CRI/CRA pelo codigo, titulo publico pelo tipo,
    LF/CDB pelo tipo + grupo, fundos pelo nome sem palavras genericas."""
    a = g.ativo.fillna("")
    cod = g.codigo.fillna("").str.strip()
    fundo = g.casa.isin(FUNDOS)
    ticker = cod.str.fullmatch(r"[A-Z]{4}[A-Z0-9]{1,4}") & ~fundo
    cri = a.str.extract(r"\b(CR[AI]):([A-Z0-9]+)")
    isin = a.str.extract(r"\b(BR[A-Z0-9]{10})\b")[0]
    curto = a.str.replace(GENERICO, " ", regex=True).str.replace(r"\s+", " ", regex=True).str.strip(" -").str.slice(0, 38)
    base = np.select([ticker, tp.notna(), cri[0].notna(), isin.notna(), fundo],
                     [cod, tp.fillna(""), cri[0].fillna("") + " " + cri[1].fillna(""), isin.fillna(""), curto],
                     g.casa + " " + g.grupo.fillna("").str.slice(0, 26))
    cx = g.casa.eq("Caixa")
    base = np.where(cx, np.where(g.categoria.eq("Operações Compromissadas"), "Compromissada " + tp.fillna(""), g.categoria), base)
    v = venc.dt.strftime("%m/%Y")
    return pd.Series(base, index=g.index).str.strip() + np.where(v.notna() & ~fundo & ~cx, " - " + v.fillna(""), "")


def implicito(df, ctx):
    """Spread implicito pelo preco (ativos sem taxa publica: cotas de FIDC, CRI sem taxa, bonds...): mediana dos ultimos
    6 retornos mensais do preco unitario acima do CDI, anualizada. Em dolar: tira o cambio e compara com o Treasury 1 ano."""
    k = lt.contribuicao(df)
    if not len(k):
        return pd.Series(dtype=float)
    k = k[(k.qt > 0) & (k.qt_ant > 0) & (k.v_ant > 0)].copy()
    rp = (k.v / k.qt) / (k.v_ant / k.qt_ant) - 1
    usd = k.chave.map(df.drop_duplicates("chave").set_index("chave").casa).isin(USD)
    e = (1 + rp) / (1 + k.mes.map(ctx.cdi_m)) - 1
    e = e.where(~usd, (1 + rp) / ((1 + k.mes.map(ctx.fx_m).fillna(0)) * (1 + k.mes.map(ctx.ust_m).fillna(0))) - 1)
    k["e"] = e.where(e.abs() < 0.2)
    k = k.sort_values(["chave", "mes"])
    k["med"] = k.groupby("chave").e.transform(lambda s: s.rolling(6, min_periods=1).median())
    return ((1 + k.set_index(["mes", "chave"]).med) ** 12 - 1) * 100


def mede(g, m, ctx, recentes, impl, ant=None):
    """Spread CDI (% a.a., equivalente), duration (anos), PU, indexador e fonte de TODAS as linhas da carteira do mes.
    Preferencia: ANBIMA (mercado) > SND + preco do fundo (mercado) > taxa contratada (CVM) > implicito pelo preco >
    media da categoria. Renda variavel (acoes, FII, FIP, ETF, FIAGRO): nao se aplica."""
    g = g.copy()
    fm, cdi, n = fim_mes(m), ctx.cdi_aa(m), len(g)
    sp, du = np.full(n, np.nan), np.full(n, np.nan)
    fonte, ind = np.full(n, "", dtype=object), np.full(n, "", dtype=object)
    cod = g.codigo.fillna("").values
    snd = ctx.snd.reindex(cod)
    venc = g.venc.fillna(pd.Series(snd.venc.values, index=g.index))
    anos = ((venc - fm).dt.days / 365.25).clip(lower=1 / 365).values
    tem_venc = venc.notna().values
    pu_f = (g.valor_veiculo / g.quantidade).where(g.quantidade > 0).values
    ca = g.casa.values.astype(str)
    nome = g.ativo.fillna("").map(lt._sem_acento).values.astype(str)
    curva = lambda t, a: lt.curva(ctx.curvas, t, fm, np.nan_to_num(a, nan=3))
    valido = lambda v: np.isfinite(v) & (v > -5) & (v < 30)          # spread fora disso = dado ruim: tenta a proxima fonte

    def poe(msk, s, d, f, i):                    # preenche so onde ainda nao ha spread
        msk = msk & np.isnan(sp)
        b = lambda v: np.broadcast_to(np.asarray(v, dtype=object if isinstance(v, str) or (np.ndim(v) and np.asarray(v).dtype == object) else float), (n,))
        sp[msk], du[msk], fonte[msk], ind[msk] = b(s)[msk], b(d)[msk], b(f)[msk], b(i)[msk]

    # 1) caixa e titulos publicos
    poe((ca == "Caixa") | ((ca == "ETF") & (np.char.find(nome, "SELIC") >= 0)), 0.0, 0.0, "caixa", "DI")
    tp = g.tp.values.astype(str)
    m_tp = ca == "Títulos Públicos"
    if m_tp.any():
        y = np.where(tp == "NTN-B", curva("IPCA", anos), curva("PRE", anos))
        _, d_cup = taxa_e_duration(anos, 0.5, 0, 0, 0, y, np.nan)
        poe(m_tp, 0.0, np.where(np.isin(tp, ["NTN-B", "NTN-F"]), d_cup, anos), "Tesouro (curva)",
            np.where(tp == "NTN-B", "IPCA", np.where(np.isin(tp, ["LTN", "NTN-F"]), "PRE", "DI")).astype(object))
    # 2) debentures: ANBIMA do dia mais proximo do fim do mes
    dia, an = ctx.anbima_mes(m, recentes)
    deb = np.isin(ca, ["DEB", "DEBIN"])
    if an is not None and deb.any():
        a = an.reindex(cod)
        tx, tipo, dur = a.taxa_ind.values.astype(float), a.tipo.values.astype(str), a.duration_anos.values.astype(float)
        real = np.isin(tipo, ["IPCA +", "IGP-M +"])
        s = np.where(tipo == "DI +", tx, np.where(tipo == "% do DI", (tx - 100) / 100 * cdi, np.nan))
        s = np.where(real, tx - lt.curva(ctx.curvas, "IPCA", pd.Timestamp(dia), np.nan_to_num(dur, nan=3)), s)
        s = np.where(tipo == "Outros", tx - lt.curva(ctx.curvas, "PRE", pd.Timestamp(dia), np.nan_to_num(dur, nan=2)), s)
        poe(deb & np.isfinite(s), s, dur, "ANBIMA", np.where(real, "IPCA", np.where(tipo == "Outros", "PRE", "DI")).astype(object))
    # 3) debentures fora da ANBIMA: taxa de emissao (SND) ajustada pelo preco do fundo / PU par do SND
    idx = snd.indice.fillna("").values.astype(str)
    di, real, pre = idx == "DI", np.isin(idx, ["IPCA", "IGP-M", "INPC"]), idx == "PRE"
    m3 = deb & np.isnan(sp) & (di | real | pre) & tem_venc
    if m3.any():
        pct, jur = snd.pct.values.astype(float), snd.juros.values.astype(float)
        s0 = np.where(di & (pct > 100) & (np.nan_to_num(jur) == 0), (pct - 100) / 100 * cdi, jur)
        s0 = np.where((np.nan_to_num(jur) == 0) & ~(di & (pct > 100)), np.nan, s0)     # "nao padrao" no SND: sem taxa util
        par = ctx.pu_par(m).reindex(cod).values.astype(float) * np.where(real, 1 + ctx.proj_ipca(m), 1)
        pp_ = pu_f / par
        pp_ = np.where((pp_ > 0.7) & (pp_ < 1.15), pp_, np.nan)
        car = ((snd.carencia - fm).dt.days / 365.25).values.astype(float)
        r = np.where(m3 & np.isfinite(s0))[0]
        s, d = np.full(n, np.nan), np.full(n, np.nan)
        s[r], d[r] = ctx.taxa(m, cod[r], anos[r], snd.cada_juros.values.astype(float)[r] / 12,
                              np.nan_to_num(snd.cada_amort.values.astype(float)[r]) / 12, car[r], np.where(di[r], cdi, 0), s0[r], pp_[r])
        s = np.where(np.abs(s - s0) > 15, s0, s)              # preco sem sentido (evento, erro de unidade): fica a taxa de emissao
        spr = np.where(di, s, np.where(real, s - curva("IPCA", d), s - curva("PRE", d)))
        poe(m3 & valido(spr), spr, d, np.where(np.isfinite(pp_), "SND + preço do fundo", "SND (taxa de emissão)").astype(object),
            np.where(di, "DI", np.where(real, "IPCA", "PRE")).astype(object))
    # 4) taxa contratada informada a CVM (LF, CDB, CRA, NC, debenture sem codigo...)
    ix = g.indexador.fillna("").values.astype(str)
    cup = g.cupom.values.astype(float)
    cup = np.where(np.nan_to_num(cup) != 0, cup, np.nan)
    pct_i, tpre = g.pct_indexador.values.astype(float), g.taxa_pre.values.astype(float)
    i_di = (np.char.find(ix, "DI de um dia") >= 0) | (np.char.find(ix, "Selic") >= 0)
    i_ip, i_pr = np.char.find(ix, "IPCA") >= 0, np.char.find(ix, "prefixada") >= 0
    m4 = np.isnan(sp) & (i_di | i_ip | i_pr) & ~np.isin(ca, SEM_SPREAD) & tem_venc
    if m4.any():
        cup_di = np.where(cup >= 50, (cup - 100) / 100 * cdi, cup)                  # % do CDI informado no campo de cupom
        taxa = np.where(i_di, np.where(np.isfinite(cup), cup_di, (pct_i - 100) / 100 * cdi), np.where(i_ip, cup, np.where(np.isfinite(tpre), tpre, cup)))
        zero = np.isin(ca, ["LF", "LFSN", "CDB", "DPGE", "LCA", "NC", "CCB", "LCI/LH"])   # pagam tudo no vencimento
        _, d = taxa_e_duration(anos, np.where(zero, anos + 1, 0.5), 0, 0, np.where(i_di, cdi, 0), np.nan_to_num(taxa), np.nan)
        spr = np.where(i_di, taxa, np.where(i_ip, taxa - curva("IPCA", d), taxa - curva("PRE", d)))
        poe(m4 & valido(spr), spr, d, "taxa contratada (CVM)", np.where(i_di, "DI", np.where(i_ip, "IPCA", "PRE")).astype(object))
    # 5) sem taxa publica: implicito pelo preco (cotas de FIDC, CRI sem taxa, bonds, outros)
    fid = ca == "FIDC"
    d_fidc = np.full(n, np.nan)
    d_fidc[fid] = [ctx.fidc_anos(c, m) for c in cod[fid]]
    m_f = fid & np.isnan(sp)
    if m_f.any():
        v = np.full(n, np.nan)
        v[m_f] = [ctx.fidc_spread(c, m, p) for c, p in zip(cod[m_f], pu_f[m_f])]
        poe(m_f & np.isfinite(v) & (v > -30) & (v < 40), v, d_fidc, "FIDC: rentabilidade da série (CVM)", "DI")
    m5 = np.isnan(sp) & ~np.isin(ca, SEM_SPREAD)
    if m5.any() and len(impl):
        v = impl.reindex(pd.MultiIndex.from_arrays([[m] * n, g.chave.values])).values.astype(float)
        usd = np.isin(ca, USD)
        d = d_fidc.copy()
        _, d_cri = taxa_e_duration(anos, 1 / 12, 1 / 12, 0, cdi, np.nan_to_num(v), np.nan)       # CRI: amortiza todo mes
        d = np.where((ca == "CRI") & tem_venc, d_cri, d)
        ust = lt.curva(ctx.ust, "UST", fm, np.nan_to_num(anos, nan=5)) if ctx.ust else np.full(n, 4.5)
        _, d_b = taxa_e_duration(anos, 0.5, 0, 0, ust, np.nan_to_num(v), np.nan)
        d = np.where((ca == "Bonds") & tem_venc, d_b, d)
        zero = np.isin(ca, ["LF", "LFSN", "CDB", "DPGE", "LCA", "NC", "CCB", "LCI/LH"])
        _, d_g = taxa_e_duration(anos, np.where(zero, anos + 1, 0.5), 0, 0, cdi, np.nan_to_num(v), np.nan)
        d = np.where(np.isnan(d) & tem_venc, d_g, d)
        poe(m5 & valido(v), v, d, np.where(usd, "implícito (preço em US$)", np.where(fid & np.isfinite(d), "implícito (cota) + prazo FIDC",
                                                                                        "implícito (preço)")).astype(object),
            np.where(usd, "USD", "DI").astype(object))
    # 6) renda variavel: nao se aplica
    na = np.isin(ca, SEM_SPREAD) & np.isnan(sp)
    fonte[na] = np.where(ca[na] == "Confidencial", "confidencial (CVM)", "não se aplica")
    # 7) o que faltar: mesmo emissor (neste fundo ou em outro fundo no mes); senao media da categoria; senao do fundo
    w = g.perc_pl.values.astype(float)
    gr = g.grupo.fillna("").values.astype(str)
    ok = np.isfinite(sp) & np.isfinite(du) & ~na & (w > 0) & ~np.isin(ca, ["Caixa", "Títulos Públicos"])         & ~np.char.startswith(fonte.astype(str), "estimado")
    for nome_g in np.unique(gr[ok]):                                    # alimenta o mapa de emissores do mes
        x = ok & (gr == nome_g)
        a = ctx.emissores.setdefault((m, nome_g), [0.0, 0.0, 0.0])
        a[0] += (sp[x] * w[x]).sum(); a[1] += (du[x] * w[x]).sum(); a[2] += w[x].sum()
    sem = np.isnan(sp) & ~na
    for nome_g in np.unique(gr[sem]):
        a = ctx.emissores.get((m, nome_g)) if nome_g else None
        if a and a[2] > 0:
            sel = sem & (gr == nome_g)
            sp[sel], fonte[sel], ind[sel] = a[0] / a[2], "estimado (mesmo emissor)", "DI"
            du[sel] = np.where(np.isnan(du[sel]), a[1] / a[2], du[sel])
    falta_s, falta_d = np.isnan(sp) & ~na, np.isnan(du) & ~na
    if falta_s.any() or falta_d.any():
        cred = ~np.isin(ca, ["Caixa", "Títulos Públicos"]) & ~na & (w > 0)

        def media(v, msk):
            ok = msk & np.isfinite(v) & (w > 0)
            return float(np.average(v[ok], weights=w[ok])) if ok.any() else np.nan
        for c in np.unique(ca[falta_s | falta_d]):
            mc = ca == c
            for v, falta, rotulo_ in ((sp, falta_s, None), (du, falta_d, " · duration estimada")):
                sel = mc & falta
                if not sel.any():
                    continue
                x, orig = media(v, mc & ~falta), "estimado (média da categoria)"
                if np.isnan(x):
                    x, orig = media(v, cred & ~falta), "estimado (média do fundo)"
                if np.isnan(x) and ant:                  # FIC sem a carteira do fundo investido no mes
                    x, orig = ant[0 if v is sp else 1], "estimado (carteira do mês anterior)"
                v[sel] = x
                if rotulo_ is None:
                    fonte[sel], ind[sel] = orig, "DI"
                else:
                    fonte[sel & ~falta_s] = fonte[sel & ~falta_s] + rotulo_
    g["spread"], g["duration"], g["pu"], g["fonte"] = sp, du, pu_f, fonte
    g["ipca"], g["pre"] = ind == "IPCA", ind == "PRE"
    g["usd"] = np.isin(ca, USD) | (ind == "USD")
    return g, dia


def derivativos_mes(gd, m, ctx, dur_ipca, dur_pre):
    """Contratos futuros do mes: tipo, lado, contratos, prazo e nocional estimado (R$)."""
    out = []
    for r in gd.itertuples():
        tipo, venc = contrato(f"{r.codigo or ''} {r.ativo or ''} {r.tipo_ativo or ''}")
        if "mercado futuro" not in str(r.categoria).lower():   # termo/opcao/swap: sem P&L de futuro
            tipo, venc = "OUTRO", None
        if tipo in ("DI1", "DAP") and not venc:                 # sem vencimento nao da para saber o prazo (ex.: trava de curva)
            tipo = "OUTRO"
        lado = -1 if "vendida" in str(r.categoria).lower() else 1
        n = abs(r.quantidade) if pd.notna(r.quantidade) else np.nan
        if venc:
            anos = max((pd.Timestamp(venc[0], venc[1], 15 if tipo == "DAP" else 1) - fim_mes(m)).days / 365.25, 1 / 252)
        else:
            anos = dur_ipca if tipo == "DAP" else dur_pre if tipo == "DI1" else 1.0
        out.append(dict(i=r.Index, tipo=tipo, lado=lado, n=n, anos=anos, fator=r.fator, categoria=r.categoria,
                        ativo=r.ativo, perc_cda=r.perc_pl))
    return out


def valor_ponto(tipo, m, ctx):
    if tipo == "DI1":
        return 1.0
    if tipo == "DAP":
        x = ctx.ipca[ctx.ipca.index <= m]
        return 0.00025 * float(x.iloc[-1]) if len(x) else np.nan
    return np.nan


def pu(tipo, m, anos, ctx):
    t = "PRE" if tipo == "DI1" else "IPCA"
    y = float(lt.curva(ctx.curvas, t, fim_mes(m), np.array([anos]))[0])
    return 1e5 / (1 + y / 100) ** anos


def nocional(c, m, ctx, pl_fundo):
    if c["tipo"] in ("DI1", "DAP") and pd.notna(c["n"]):
        return c["lado"] * c["n"] * pu(c["tipo"], m, c["anos"], ctx) * valor_ponto(c["tipo"], m, ctx)
    if c["tipo"] in ("DOL", "WDO") and pd.notna(c["n"]):
        return c["lado"] * c["n"] * (50000 if c["tipo"] == "DOL" else 10000) * ctx.ptax_em(m)
    return c["perc_cda"] / 100 * pl_fundo if pd.notna(c["perc_cda"]) else np.nan


def pnl_futuro(c, p, m, ctx):
    """P&L estimado (R$) do contrato entre o fim do mes p e o fim do mes m, com a posicao de p."""
    if pd.isna(c["n"]):
        return 0.0
    cdi = float(ctx.cdi_m.get(m, 0.0))
    if c["tipo"] in ("DI1", "DAP"):
        a1 = max(c["anos"] - (fim_mes(m) - fim_mes(p)).days / 365.25, 1 / 252)
        pu0, pu1 = pu(c["tipo"], p, c["anos"], ctx), pu(c["tipo"], m, a1, ctx)
        if c["tipo"] == "DI1":
            return c["lado"] * c["n"] * (pu1 - pu0 * (1 + cdi))
        i0, i1 = ctx.ipca[ctx.ipca.index <= p], ctx.ipca[ctx.ipca.index <= m]
        inf = float(i1.iloc[-1] / i0.iloc[-1] - 1) if len(i0) and len(i1) else 0.0
        return c["lado"] * c["n"] * valor_ponto("DAP", m, ctx) * (pu1 - pu0 * (1 + cdi) / (1 + inf))
    if c["tipo"] in ("DOL", "WDO"):
        return c["lado"] * c["n"] * (50000 if c["tipo"] == "DOL" else 10000) * (ctx.ptax_em(m) - ctx.ptax_em(p))
    return 0.0


def detalhe(nome, cnpj, df, ctx, pl, rent):
    """Tudo que as abas Carteira e Retorno mostram para um fundo + tabela de ativos (busca/top/grupos) de cada mes."""
    df = df.copy()
    # fecha 100% do PL em todo mes: a diferenca vira uma linha explicita (confidencial se o mes tem parte relevante assim)
    nd = df[~df.derivativo]
    falta = (100 - nd.groupby("mes").perc_pl.sum()).loc[lambda x: x.abs() > 0.05]
    if len(falta):
        conf = nd[nd.confidencial].groupby("mes").perc_pl.sum().reindex(falta.index).fillna(0) >= 5
        df = pd.concat([df, pd.DataFrame({"fundo": nome, "mes": falta.index, "categoria": lt.AJUSTE, "perc_pl": falta.values,
                                          "ativo": "Diferença entre PL e carteira informada (CVM)", "derivativo": False,
                                          "confidencial": conf.values, "chave": "ajuste|" + falta.index})], ignore_index=True)
    for c in ("emissor", "codigo", "ativo", "tipo_ativo", "indexador", "categoria"):     # coluna toda vazia vira texto
        if c in df.columns:
            df[c] = df[c].astype(object)
    if len(ctx.snd):                                    # emissor faltando: nome do emissor pelo SND / ANBIMA
        df["emissor"] = df.emissor.fillna(df.codigo.map(ctx.snd.empresa))
    if ctx.an_h is not None and len(ctx.an_h):
        df["emissor"] = df.emissor.fillna(df.codigo.map(ctx.an_h.drop_duplicates("codigo", keep="last").set_index("codigo").nome))
    lixo = r"(?i)\bBONDS?\s*\d*\s*DAYS?\s*SETTLE\b|\s[\d.,]{5,}\b"           # bonds: 'BONDS 2 DAYS SETTLE - emissor - qtd'
    df["emissor"] = df.emissor.astype(object).str.replace(lixo, " ", regex=True).str.strip(" -").replace("", "Emissor não informado")
    emis = df[["emissor", "categoria"]].drop_duplicates()
    k = lambda e: e if isinstance(e, str) else None                            # NaN nao serve de chave
    gmap = {(k(e), c): lt.grupo(e, c) for e, c in zip(emis.emissor, emis.categoria)}
    df["grupo"] = [gmap[(k(e), c)] for e, c in zip(df.emissor, df.categoria)]
    curto = df.ativo.str.replace(GENERICO, " ", regex=True).str.replace(r"\s+", " ", regex=True).str.strip(" -").str.slice(0, 40)
    df["grupo"] = df.grupo.fillna(curto)
    df["casa"] = casa(df, ctx)
    df["venc"] = vencimento(df)
    df["tp"] = tipo_tp(df)
    df["rotulo"] = rotulo(df, df.venc.fillna(df.codigo.map(ctx.snd.venc) if len(ctx.snd) else df.venc), df.tp)
    impl = implicito(df, ctx)
    meses = sorted(df.mes.unique())
    recentes = set(meses[-4:])
    pl_f = {m: float(pl.get((m, cnpj), np.nan)) for m in meses}
    cart_meses, ricos, tab_ativos, ant = [], {}, [], None
    for m in meses:
        g = df[df.mes == m]
        e, dia = mede(g[~g.derivativo], m, ctx, recentes, impl, ant)
        ricos[m] = (e, g[g.derivativo])
        e["fin"] = e.perc_pl / 100 * pl_f[m] / 1e6
        cats = [{"categoria": c, "perc": x.perc_pl.sum(), "fin": x.fin.sum(), "spread": wavg(x.spread, x.perc_pl),
                 "duration": wavg(x.duration, x.perc_pl)} for c, x in e.groupby("casa")]
        cats.sort(key=lambda r: -r["perc"])
        total = {"categoria": "Total", "perc": e.perc_pl.sum(), "fin": e.fin.sum(), "spread": wavg(e.spread, e.perc_pl),
                 "duration": wavg(e.duration, e.perc_pl)}
        pos = e[e.perc_pl > 0]
        fontes = (pos.groupby(pos.fonte.str.split(" · ").str[0]).perc_pl.sum() / max(pos.perc_pl.sum(), 1e-9) * 100).round(1)
        if total["spread"] is not None and total["duration"] is not None:
            ant = (total["spread"], total["duration"])
        cart_meses.append({"mes": m, "pl_mm": pl_f[m] / 1e6, "anbima": dia, "conf": e[e.confidencial].perc_pl.sum(),
                           "cats": cats, "total": total, "fontes": fontes.sort_values(ascending=False).to_dict()})
        # tabela de ativos do mes (rotulo curto), maior posicao primeiro
        x = e[e.perc_pl != 0].sort_values("perc_pl", ascending=False)
        x = x.assign(sw=(x.spread * x.perc_pl).fillna(0), sp_=x.perc_pl.where(x.spread.notna(), 0),
                     dw=(x.duration * x.perc_pl).fillna(0), dp=x.perc_pl.where(x.duration.notna(), 0))
        t = x.groupby("rotulo", sort=False).agg(categoria=("casa", "first"), grupo=("grupo", "first"), emissor=("emissor", "first"),
                                                codigo=("codigo", "first"), perc=("perc_pl", "sum"), sw=("sw", "sum"), sp_=("sp_", "sum"),
                                                dw=("dw", "sum"), dp=("dp", "sum"), pu=("pu", "first"), fonte=("fonte", "first")).reset_index()
        t["spread"] = (t.sw / t.sp_).where(t.sp_ != 0)
        t["duration"] = (t.dw / t.dp).where(t.dp != 0)
        tab_ativos.append(t.drop(columns=["sw", "sp_", "dw", "dp"]).assign(mes=m))

    # ---- atribuicao mensal (retorno por categoria, hedges alocados nos ativos protegidos)
    real = rent[rent.fundo == nome].set_index("mes").retorno
    k = lt.contribuicao(df)
    if len(k):                                   # variacao do preco unitario de cada posicao mantida (v/qt)
        k = k[(k.qt > 0) & (k.qt_ant > 0) & (k.v_ant > 0)]
        k = k.assign(rp=(k.v / k.qt) / (k.v_ant / k.qt_ant) - 1)
        preco = k.groupby(["mes", "chave"]).rp.first()
    else:
        preco = pd.Series(dtype=float)
    atr_meses, cats_atr, ativos, contrib = [], {}, {}, []
    id_ativo = lambda a, c, gr, cod, em: ativos.setdefault((a, c), {"i": len(ativos), "ativo": a, "categoria": c, "grupo": gr,
                                                                     "codigo": cod, "emissor": em})["i"]
    for j in range(1, len(meses)):
        p, m = meses[j - 1], meses[j]
        if m not in real.index or pd.isna(real.get(m)):
            continue
        e0, gd0 = ricos[p]
        w = e0[e0.perc_pl.abs() > 0].copy()            # confidenciais entram (rendem CDI + residuo): a soma sempre fecha
        cdi_m = float(ctx.cdi_m.get(m, 0.0))
        # 1) retorno de cada ativo: esperado = CDI + spread proprio; usa a variacao real de preco quando ela e critica
        #    (queda maior que 1,5 p.p. abaixo do esperado = pagamento de cupom/amortizacao -> fica o esperado)
        esperado = (1 + cdi_m) * (1 + w.spread.fillna(0).clip(-5, 30) / 100) ** (1 / 12) - 1
        # + marcacao a mercado: -duration x (variacao da curva IPCA/pre no prazo + variacao do spread de mercado)
        D = w.duration.fillna(0).clip(0, 30).values
        dc = np.zeros(len(w))
        for flag, t in ((w.ipca.values, "IPCA"), (w.pre.values, "PRE")):
            if flag.any():
                dc[flag] = lt.curva(ctx.curvas, t, fim_mes(m), D[flag]) - lt.curva(ctx.curvas, t, fim_mes(p), D[flag])
        e1 = ricos[m][0].drop_duplicates("chave").set_index("chave")
        ds = (w.chave.map(e1.spread) - w.spread).where(w.fonte.isin(MERCADO) & w.chave.map(e1.fonte).isin(MERCADO))
        y = ctx.cdi_aa(p) / 100 + w.spread.fillna(0).values / 100
        esperado = esperado - D / (1 + y) * (np.nan_to_num(dc) + ds.fillna(0).clip(-5, 5).values) / 100
        rp = w.chave.map(preco.loc[m]) if len(preco) and m in preco.index.get_level_values(0) else pd.Series(np.nan, index=w.index)
        r = rp.where(rp.notna() & (rp >= esperado - 0.015) & (rp <= esperado + 0.03), esperado)
        cx = w.casa.isin(CAIXA)
        r[cx] = cdi_m
        ativo_c = w.perc_pl * r
        # 2) derivativos -> ativos protegidos
        ip = w[w.ipca & w.perc_pl.gt(0)]
        dur_ipca = wavg(ip.duration, ip.perc_pl) or 3.0
        pr = w[w.pre & w.perc_pl.gt(0)]
        dur_pre = wavg(pr.duration, pr.perc_pl) or 1.0
        hedge = 0.0
        for c in derivativos_mes(gd0, p, ctx, dur_ipca, dur_pre):
            v = pnl_futuro(c, p, m, ctx) * c["fator"] / 100 * 100 if pd.notna(c["fator"]) else 0.0  # % do PL do fundo
            if not v or np.isnan(v):
                continue
            alvo = (w.ipca if c["tipo"] == "DAP" else w.pre if c["tipo"] == "DI1" and w.pre.any() else w.ipca
                    if c["tipo"] == "DI1" else w.usd if c["tipo"] in ("DOL", "WDO", "DDI") else pd.Series(False, index=w.index))
            alvo = alvo & w.perc_pl.gt(0)
            if not alvo.any():
                alvo = w.perc_pl.gt(0) & ~cx
            ativo_c[alvo] += v * w.perc_pl[alvo] / w.perc_pl[alvo].sum()
            hedge += v
        # 3) resto (taxas, negociacao, marcacao fora do esperado) -> todas as categorias que nao sao caixa, pelo peso
        resid = float(real[m]) * 100 - ativo_c.sum()
        cr = w.perc_pl.gt(0) & ~cx
        if not cr.any():
            cr = w.perc_pl.gt(0)
        ativo_c[cr] += resid * w.perc_pl[cr] / w.perc_pl[cr].sum()
        atr_meses.append({"mes": m, "real": float(real[m]) * 100, "cdi": float(ctx.cdi_m.get(m, np.nan)) * 100,
                          "hedge": hedge, "resid": resid})
        jm = len(atr_meses) - 1
        por_cat = ativo_c.groupby(w.casa).sum()
        peso = w.groupby("casa").perc_pl.sum()
        for c in set(por_cat.index) | set(cats_atr):
            cats_atr.setdefault(c, {"contrib": [None] * jm, "peso": [None] * jm})
            cats_atr[c]["contrib"].append(float(por_cat.get(c, 0.0)))
            cats_atr[c]["peso"].append(float(peso.get(c, 0.0)))
        nc = w[~cx & ~w.casa.eq("Confidencial")].assign(c_=ativo_c[~cx & ~w.casa.eq("Confidencial")])                               # caixa/compromissada fora do ranking
        pa = nc.groupby(["rotulo", "casa"]).agg(v=("c_", "sum"), p=("perc_pl", "sum"), gr=("grupo", "first"),
                                                cod=("codigo", "first"), em=("emissor", "first"))
        for (a, c), x in pa.iterrows():
            if abs(x.v) > 1e-5:
                contrib.append([id_ativo(a, c, x.gr, x.cod, x.em), jm, round(float(x.v), 6), round(float(x.p), 4)])
    for c in cats_atr.values():
        for key in ("contrib", "peso"):
            c[key] += [None] * (len(atr_meses) - len(c[key]))
    tab = pd.concat(tab_ativos, ignore_index=True) if tab_ativos else pd.DataFrame()
    return limpo({"nome": nome, "cnpj": cnpj, "carteira": cart_meses,
                  "retorno": {"meses": atr_meses, "categorias": cats_atr,
                              "ativos": sorted(ativos.values(), key=lambda x: x["i"]), "contrib": contrib}}), tab


# ------------------------------------------------------------------ pre-carga de todos os fundos
def pasta_dia():
    return lt.CACHE / "painel" / str(date.today())


def precarrega(log=None):
    escreve = log or PRE["log"].append
    PRE["status"] = "rodando"
    try:
        fundos = [(n, lt.cnpj_of(x), True) for n, x in lt.NOSSOS_FUNDOS.items()] + \
                 [(n, lt.cnpj_of(x), False) for n, x in lt.PEERS.items()]
        dest = pasta_dia()
        chave = "|".join(sorted(c for _, c, _ in fundos)) + "|" + lt.DESDE + f"|m{METODO}"
        if (dest / "visao.json").exists() and json.loads((dest / "visao.json").read_text(encoding="utf-8")).get("chave") == chave:
            escreve("Usando o painel já processado hoje.")
            PRE.update(status="ok", pasta=str(dest))
            return
        dest.mkdir(parents=True, exist_ok=True)
        for velho in (lt.CACHE / "painel").glob("*"):
            if velho != dest and velho.is_dir():
                for f in velho.glob("*"):
                    f.unlink(missing_ok=True)
                velho.rmdir()
        escreve(f"Baixando/atualizando dados da CVM de {len(fundos)} fundos desde {lt.DESDE}...")
        workers = max(1, min(4, (os.cpu_count() or 2)))
        meses, dia = lt.prepara(sorted(c for _, c, _ in fundos), lt.DESDE, workers=workers, log=escreve)
        ctx = Contexto(escreve, meses)
        escreve("Carregando carteiras...")
        carga = lt.carrega_pos(meses)
        pl = carga[1]
        nomes = {n: c for n, c, _ in fundos}
        d = lt.diario(nomes, dia)
        rent = lt.rentabilidade(d, ctx.cdi_m) if len(d) else pd.DataFrame(columns=["fundo", "mes", "retorno"])
        for i, (n, c, _) in enumerate(fundos, 1):
            escreve(f"[{i}/{len(fundos)}] {n}: look-through, carteira e atribuição...")
            fid = re.sub(r"\D", "", c)
            try:                                 # um fundo com problema nao derruba os outros
                df = lt.lookthrough({n: c}, meses, carga)
                if not len(df):
                    escreve(f"   {n}: a CVM não tem carteira deste CNPJ no período (só cota/retorno).")
                    continue
                det, tab = detalhe(n, c, df, ctx, pl, rent)
            except Exception as e:
                escreve(f"   {n}: pulado ({type(e).__name__}: {e})")
                continue
            (dest / (fid + ".json")).write_text(json.dumps(det), encoding="utf-8")
            tab.to_parquet(dest / (fid + "_ativos.parquet"), index=False)
            df.drop(columns=["chave", "fator"], errors="ignore").to_parquet(dest / (fid + ".parquet"), index=False)
        # visao geral: cotas diarias de todos (para retorno em qualquer periodo, consolidado e risco x retorno)
        datas = sorted(d.data.unique())
        idx = pd.DatetimeIndex(datas)
        q = d.pivot_table(index="data", columns="fundo", values="VL_QUOTA").reindex(idx)
        plf = d.pivot_table(index="data", columns="fundo", values="VL_PATRIM_LIQ").reindex(idx)
        cdi_idx = (1 + ctx.cdi_d.reindex(idx).fillna(0)).cumprod()
        cores, k_nosso = {}, 0
        for n, c, nosso in fundos:
            cores[n] = NOSSAS_CORES[k_nosso % len(NOSSAS_CORES)] if nosso else "#9aa0a6"
            k_nosso += nosso
        visao = {"chave": chave, "gerado": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M"), "desde": lt.DESDE,
                 "ult_carteira": max(meses) if meses else None, "anbima": ctx.an_dias[-1] if ctx.an_dias else None,
                 "datas": [t.strftime("%Y-%m-%d") for t in idx], "cdi": cdi_idx.tolist(),
                 "fundos": [{"nome": n, "cnpj": c, "id": re.sub(r"\D", "", c), "nosso": nosso, "cor": cores[n],
                             "tem_carteira": (dest / (re.sub(r"\D", "", c) + ".json")).exists(),
                             "cota": q[n].tolist() if n in q else [], "pl": float(plf[n].dropna().iloc[-1]) / 1e6 if n in plf and plf[n].notna().any() else None}
                            for n, c, nosso in fundos],
                 "cat_cores": CAT_CORES}
        (dest / "visao.json").write_text(json.dumps(limpo(visao)), encoding="utf-8")
        escreve("Pronto.")
        PRE.update(status="ok", pasta=str(dest))
    except Exception as e:
        import traceback
        PRE.update(status="erro", erro=f"{type(e).__name__}: {e}")
        escreve("ERRO: " + PRE["erro"])
        escreve(traceback.format_exc()[-1500:])


def _inicio():                              # servidor (pasta/Docker): pre-carga em segundo plano
    if PRE["status"] == "parado":
        threading.Thread(target=precarrega, daemon=True).start()


app.add_event_handler("startup", _inicio)


# ------------------------------------------------------------------ API
@app.middleware("http")
async def sem_cache(req, call_next):         # pagina/JS/CSS novos aparecem logo depois de atualizar o arquivo
    r = await call_next(req)
    if not req.url.path.startswith("/api/"):
        r.headers["Cache-Control"] = "no-cache"
    return r


@app.get("/api/inicial")
def inicial():
    return {"status": PRE["status"], "log": PRE["log"][-8:], "erro": PRE["erro"]}


def _arquivo(nome):
    if PRE["status"] != "ok":
        raise HTTPException(503, "painel ainda carregando")
    f = os.path.join(PRE["pasta"], nome)
    if not os.path.exists(f):
        raise HTTPException(404, "sem dados")
    return f


@app.get("/api/visao")
def visao():
    return FileResponse(_arquivo("visao.json"), media_type="application/json")


@app.get("/api/fundo/{fid}")
def fundo(fid: str):
    return FileResponse(_arquivo(re.sub(r"\D", "", fid) + ".json"), media_type="application/json")


_ATIVOS = {}


@app.get("/api/ativos/{fid}")
def ativos_mes(fid: str, mes: str):
    """Todos os ativos do fundo no mes (rotulo, categoria, grupo, % PL, spread, duration, PU, fonte)."""
    f = _arquivo(re.sub(r"\D", "", fid) + "_ativos.parquet")
    if f not in _ATIVOS:
        if len(_ATIVOS) > 30:
            _ATIVOS.clear()
        _ATIVOS[f] = pd.read_parquet(f)
    t = _ATIVOS[f]
    return JSONResponse(limpo(t[t.mes == mes].drop(columns="mes").to_dict("records")))


@app.get("/api/csv/{fid}")
def csv(fid: str):
    df = pd.read_parquet(_arquivo(re.sub(r"\D", "", fid) + ".parquet"))
    return Response(df.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="carteira_{fid}.csv"'})


@app.get("/")
def index():
    return HTMLResponse(PAGINA)


PAGINA = r"""<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Painel de fundos</title>
<link rel="icon" href="data:image/svg+xml,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 32 32'><rect x='4' y='16' width='5' height='12' rx='1' fill='%232a78d6'/><rect x='13' y='9' width='5' height='19' rx='1' fill='%231baf7a'/><rect x='22' y='4' width='5' height='24' rx='1' fill='%23eb6834'/></svg>">
<style>
:root {
  --bg: #f5f6f8; --surface: #ffffff; --surface-2: #f0f1f4; --line: #e3e5ea; --line-2: #d0d3da;
  --text-1: #14161a; --text-2: #4d525c; --muted: #8a8f99;
  --accent: #2a78d6; --accent-soft: rgba(42,120,214,.10); --nosso: #eb6834; --nosso-soft: rgba(235,104,52,.12);
  --good: #0f7b3f; --bad: #c43d3d; --radius: 12px;
  --shadow: 0 1px 2px rgba(16,24,40,.05), 0 1px 3px rgba(16,24,40,.06);
  --font: "Segoe UI", -apple-system, BlinkMacSystemFont, Roboto, "Helvetica Neue", Arial, sans-serif;
}
@media (prefers-color-scheme: dark) {
  :root:not([data-theme="light"]) {
    --bg: #0f1012; --surface: #17191c; --surface-2: #1f2226; --line: #2a2d32; --line-2: #3a3e45;
    --text-1: #f2f3f5; --text-2: #b7bcc6; --muted: #858b96; --accent: #3987e5; --accent-soft: rgba(57,135,229,.16);
    --good: #2fbf6b; --bad: #ef6b6b; --shadow: none; --nosso-soft: rgba(235,104,52,.18);
  }
}
:root[data-theme="dark"] {
  --bg: #0f1012; --surface: #17191c; --surface-2: #1f2226; --line: #2a2d32; --line-2: #3a3e45;
  --text-1: #f2f3f5; --text-2: #b7bcc6; --muted: #858b96; --accent: #3987e5; --accent-soft: rgba(57,135,229,.16);
  --good: #2fbf6b; --bad: #ef6b6b; --shadow: none; --nosso-soft: rgba(235,104,52,.18);
}
* { box-sizing: border-box; }
html, body { margin: 0; background: var(--bg); color: var(--text-1); font: 14px/1.5 var(--font); }
h1 { font-size: 16px; margin: 0; }
h2 { font-size: 20px; margin: 0; letter-spacing: -.01em; }
h3 { font-size: 14px; margin: 0 0 2px; }
button { font: inherit; cursor: pointer; }
input, select { font: inherit; color: var(--text-1); background: var(--surface); border: 1px solid var(--line-2); border-radius: 8px; padding: 8px 10px; }

.barra { position: sticky; top: 0; z-index: 20; display: flex; align-items: center; gap: 18px; padding: 8px 22px; background: var(--surface); border-bottom: 1px solid var(--line); }
.marca { display: flex; align-items: center; gap: 10px; min-width: 0; }
.logo { width: 26px; height: 26px; flex: none; } .logo.grande { width: 48px; height: 48px; }
.datas { margin: 0; font-size: 11.5px; color: var(--text-2); white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
.abas { display: flex; gap: 4px; margin-left: auto; }
.abas button { border: 0; background: transparent; color: var(--text-2); padding: 8px 14px; border-radius: 8px; font-weight: 600; }
.abas button:hover { background: var(--surface-2); }
.abas button[aria-selected="true"] { background: var(--accent-soft); color: var(--accent); }
.botao { border: 1px solid var(--line-2); background: var(--surface); color: var(--text-1); border-radius: 8px; padding: 7px 12px; }
.botao:hover { background: var(--surface-2); }
.botao.icone { width: 36px; padding: 7px 0; }
.botao.sel { border-color: var(--accent); color: var(--accent); background: var(--accent-soft); font-weight: 600; }

main { max-width: 1380px; margin: 0 auto; padding: 18px 24px 60px; }
.cab { display: flex; flex-wrap: wrap; gap: 14px; align-items: flex-end; justify-content: space-between; margin-bottom: 14px; }
.titulo-fundo { min-width: 0; }
.titulo-fundo h2 { display: flex; align-items: center; gap: 8px; }
.tag { font-size: 11px; font-weight: 700; color: var(--nosso); background: var(--nosso-soft); border-radius: 6px; padding: 2px 7px; letter-spacing: .02em; }
.sub { color: var(--muted); font-size: 12px; margin: 2px 0 0; }
.controles { display: flex; flex-wrap: wrap; gap: 8px; align-items: center; }
.controles label { font-size: 12px; color: var(--text-2); display: flex; gap: 6px; align-items: center; }

/* busca de fundo por digitacao */
.busca { position: relative; width: min(420px, 100%); }
.busca input { width: 100%; padding-left: 34px; }
.busca::before { content: "⌕"; position: absolute; left: 11px; top: 6px; font-size: 17px; color: var(--muted); }
.sugestoes { position: absolute; z-index: 30; left: 0; right: 0; top: 100%; margin: 4px 0 0; padding: 4px; list-style: none; max-height: 340px; overflow: auto;
  background: var(--surface); border: 1px solid var(--line-2); border-radius: 10px; box-shadow: 0 10px 28px rgba(0,0,0,.16); }
.sugestoes li { padding: 7px 9px; border-radius: 7px; cursor: pointer; display: flex; justify-content: space-between; gap: 10px; font-size: 13px; }
.sugestoes li small { color: var(--muted); }
.sugestoes li:hover, .sugestoes li.ativo { background: var(--surface-2); }

.grade { display: grid; gap: 16px; margin-bottom: 16px; }
.g2 { grid-template-columns: repeat(2, minmax(0, 1fr)); }
.g-tabela { grid-template-columns: minmax(0, 3fr) minmax(0, 2fr); }
.cartao { background: var(--surface); border: 1px solid var(--line); border-radius: var(--radius); box-shadow: var(--shadow); padding: 14px 16px; min-width: 0; margin-bottom: 16px; }
.grade > .cartao { margin-bottom: 0; }
.cartao .sub { margin-bottom: 8px; }
.kpis { display: grid; grid-template-columns: repeat(auto-fit, minmax(150px, 1fr)); gap: 12px; margin-bottom: 16px; }
.kpi { background: var(--surface); border: 1px solid var(--line); border-radius: var(--radius); box-shadow: var(--shadow); padding: 12px 14px; }
.kpi .r { font-size: 12px; color: var(--text-2); }
.kpi .v { font-size: 24px; font-weight: 700; letter-spacing: -.02em; margin-top: 2px; font-variant-numeric: tabular-nums; }
.kpi .s { font-size: 11.5px; color: var(--muted); }
.pos { color: var(--good); } .neg { color: var(--bad); }

.tabela-wrap { overflow: auto; border: 1px solid var(--line); border-radius: 10px; }
table { border-collapse: collapse; width: 100%; font-size: 13px; }
th, td { padding: 7px 11px; border-bottom: 1px solid var(--line); text-align: left; white-space: nowrap; }
th { position: sticky; top: 0; background: var(--surface-2); color: var(--text-2); font-weight: 600; font-size: 12px; z-index: 1; }
th.ord { cursor: pointer; } th.ord:hover { color: var(--text-1); }
tbody tr:hover td { background: var(--accent-soft); }
td.n, th.n { text-align: right; font-variant-numeric: tabular-nums; }
td.texto { white-space: normal; min-width: 200px; }
tr.total td { font-weight: 700; border-top: 2px solid var(--line-2); background: var(--surface-2); }
tr.secao td { font-size: 11.5px; text-transform: uppercase; letter-spacing: .05em; color: var(--muted); background: var(--surface); padding-top: 14px; }
tr.nosso td:first-child { box-shadow: inset 3px 0 0 var(--nosso); }
.ponto { display: inline-block; width: 9px; height: 9px; border-radius: 3px; margin-right: 6px; vertical-align: 0; }
.nota { color: var(--muted); font-size: 11.5px; margin: 8px 0 0; }

/* filtros de ativos: categorias liga/desliga + busca + quantos */
.filtros { display: flex; flex-wrap: wrap; gap: 10px 16px; align-items: center; justify-content: space-between; }
.chips { display: flex; flex-wrap: wrap; gap: 6px; }
.chip-t { border: 1px solid var(--line-2); background: var(--surface); color: var(--text-2); border-radius: 999px; padding: 4px 11px; font-size: 12.5px; }
.chip-t[aria-pressed="true"] { border-color: var(--accent); background: var(--accent-soft); color: var(--accent); font-weight: 600; }
.chip-t.mudo { border-style: dashed; }
.filtros .busca { width: min(320px, 100%); }
.muted { color: var(--muted); font-size: 12px; }

/* selecao multipla (consolidado) */
.multi { display: flex; flex-wrap: wrap; gap: 6px; align-items: center; padding: 6px; border: 1px solid var(--line-2); border-radius: 10px; background: var(--surface); min-height: 44px; }
.multi .chip { display: inline-flex; align-items: center; gap: 5px; border-radius: 999px; padding: 3px 4px 3px 10px; font-size: 12.5px; background: var(--surface-2); }
.multi .chip.nosso { background: var(--nosso-soft); }
.multi .chip button { border: 0; background: none; color: var(--muted); font-size: 15px; line-height: 1; padding: 0 4px; }
.multi .busca { width: 220px; flex: 1; }
.multi .busca input { border: 0; padding: 5px 5px 5px 30px; }
.multi .busca::before { top: 2px; }

.carregando { position: fixed; inset: 0; z-index: 50; display: grid; place-items: center; background: var(--bg); }
.carregando[hidden] { display: none; }
.carregando .caixa { width: min(560px, 92vw); text-align: center; }
.carregando h2 { margin-top: 12px; }
.descricao { color: var(--text-2); font-size: 13px; }
.barra-prog { height: 4px; background: var(--line); border-radius: 4px; overflow: hidden; margin: 18px 0 10px; }
.barra-prog div { height: 100%; width: 35%; background: var(--accent); border-radius: 4px; animation: vai 1.4s ease-in-out infinite; }
@keyframes vai { 0% { transform: translateX(-100%); } 100% { transform: translateX(300%); } }
pre { text-align: left; white-space: pre-wrap; font: 12px/1.5 ui-monospace, Consolas, monospace; color: var(--text-2); background: var(--surface);
  border: 1px solid var(--line); border-radius: 8px; padding: 10px; max-height: 220px; overflow: auto; margin: 0; }
.vazio { padding: 60px 0; text-align: center; color: var(--text-2); }

@media (max-width: 980px) {
  html, body { overflow-x: hidden; }
  .barra { flex-wrap: wrap; padding: 8px 14px; gap: 8px; }
  .abas { margin-left: 0; order: 3; width: 100%; }
  main { padding: 14px 14px 50px; }
  .g2, .g-tabela { grid-template-columns: minmax(0, 1fr); }
  .kpi .v { font-size: 20px; }
}

</style>
<script src="https://cdn.jsdelivr.net/npm/plotly.js-dist-min@2.35.2/plotly.min.js"></script>
</head>
<body>
<header class="barra">
  <div class="marca">
    <svg class="logo" viewBox="0 0 32 32" aria-hidden="true"><rect x="4" y="16" width="5" height="12" rx="1.5" fill="#2a78d6"/><rect x="13.5" y="9" width="5" height="19" rx="1.5" fill="#1baf7a"/><rect x="23" y="4" width="5" height="24" rx="1.5" fill="#eb6834"/></svg>
    <div><h1>Painel de fundos</h1><p id="datas" class="datas"></p></div>
  </div>
  <nav class="abas" role="tablist">
    <button role="tab" data-aba="carteira" aria-selected="true">Carteira</button>
    <button role="tab" data-aba="retorno" aria-selected="false">Retorno</button>
    <button role="tab" data-aba="consolidado" aria-selected="false">Consolidado</button>
  </nav>
  <button id="tema" class="botao icone" title="Tema claro/escuro" aria-label="Alternar tema">◐</button>
</header>
<main id="conteudo"></main>

<div id="carregando" class="carregando">
  <div class="caixa">
    <svg class="logo grande" viewBox="0 0 32 32" aria-hidden="true"><rect x="4" y="16" width="5" height="12" rx="1.5" fill="#2a78d6"/><rect x="13.5" y="9" width="5" height="19" rx="1.5" fill="#1baf7a"/><rect x="23" y="4" width="5" height="24" rx="1.5" fill="#eb6834"/></svg>
    <h2>Preparando o painel</h2>
    <p class="descricao">Baixando e processando os dados de todos os fundos. Na primeira vez do dia pode levar alguns minutos.</p>
    <div class="barra-prog"><div></div></div>
    <pre id="log-carga"></pre>
  </div>
</div>
<script>
// Painel de fundos - 3 abas: Carteira, Retorno, Consolidado. Dados processados pelo servidor (/api).
const $ = (s, el = document) => el.querySelector(s);
const nf = (d) => new Intl.NumberFormat("pt-BR", { minimumFractionDigits: d, maximumFractionDigits: d });
const num = (v, d = 2) => (v == null || !isFinite(v) ? "–" : nf(d).format(v));
const pct = (v, d = 2) => (v == null || !isFinite(v) ? "–" : nf(d).format(v) + "%");
const pp = (v, d = 2) => (v == null || !isFinite(v) ? "–" : (v > 0 ? "+" : "") + nf(d).format(v) + " pp");
const cls = (v) => (v == null ? "" : v >= 0 ? "pos" : "neg");
const esc = (s) => String(s ?? "").replace(/[&<>"]/g, (c) => ({ "&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;" }[c]));
const css = (v) => getComputedStyle(document.documentElement).getPropertyValue(v).trim();
const MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"];
const mesBR = (m) => (m ? MESES[+m.slice(5, 7) - 1] + "/" + m.slice(0, 4) : "–");
const dataBR = (d) => (d ? d.slice(8, 10) + "/" + d.slice(5, 7) + "/" + d.slice(0, 4) : "–");
const semAcento = (t) => String(t).normalize("NFD").replace(/[̀-ͯ]/g, "").toUpperCase();
const espera = (ms) => new Promise((ok) => setTimeout(ok, ms));
const st = { V: null, F: {}, aba: "carteira", fundo: null, mesCart: {}, ret: { de: null, ate: null }, cons: { sel: null, de: null, ate: null, ord: "ret", asc: false },
  A: {}, cf: { fora: new Set(["Caixa", "Confidencial"]), n: 15, q: "" }, rf: { fora: new Set(["Caixa", "Confidencial"]), n: 15, q: "" } };

// ------------------------------------------------------------------ tema e abas
try { const t = localStorage.getItem("tema"); if (t) document.documentElement.dataset.theme = t; } catch (e) {}
$("#tema").onclick = () => {
  const escuro = document.documentElement.dataset.theme === "dark" || (!document.documentElement.dataset.theme && matchMedia("(prefers-color-scheme: dark)").matches);
  document.documentElement.dataset.theme = escuro ? "light" : "dark";
  try { localStorage.setItem("tema", document.documentElement.dataset.theme); } catch (e) {}
  if (st.V) desenha();
};
document.querySelectorAll(".abas button").forEach((b) => (b.onclick = () => {
  st.aba = b.dataset.aba;
  document.querySelectorAll(".abas button").forEach((x) => x.setAttribute("aria-selected", x === b));
  desenha();
}));

// ------------------------------------------------------------------ carga
(async function inicia() {
  for (;;) {
    let s;
    try { s = await (await fetch("/api/inicial")).json(); } catch (e) { await espera(2000); continue; }
    $("#log-carga").textContent = (s.log || []).join("\n") || "Iniciando...";
    if (s.status === "erro") return ($("#log-carga").textContent += "\n\nErro: " + s.erro);
    if (s.status === "ok") break;
    await espera(2000);
  }
  st.V = await (await fetch("/api/visao")).json();
  const V = st.V;
  st.fundo = (V.fundos.find((f) => f.nosso && f.tem_carteira) || V.fundos.find((f) => f.tem_carteira) || V.fundos[0]).id;
  st.cons.sel = new Set(V.fundos.map((f) => f.id));
  $("#datas").textContent = `${V.fundos.length} fundos (${V.fundos.filter((f) => f.nosso).length} nossos) · carteiras até ${mesBR(V.ult_carteira)} · cotas até ${dataBR(V.datas[V.datas.length - 1])} · ANBIMA ${dataBR(V.anbima)}`;
  $("#carregando").hidden = true;
  desenha();
})();
async function detalhe(id) {
  if (!st.F[id]) st.F[id] = await (await fetch("/api/fundo/" + id)).json();
  return st.F[id];
}
const fundoPor = (id) => st.V.fundos.find((f) => f.id === id);

// ------------------------------------------------------------------ busca de fundo (digitando)
function campoBusca(onPick, placeholder = "Digite o nome ou CNPJ do fundo...", soComCarteira = false) {
  const wrap = document.createElement("div");
  wrap.className = "busca";
  wrap.innerHTML = `<input type="search" placeholder="${placeholder}" autocomplete="off"><ul class="sugestoes" hidden></ul>`;
  const inp = $("input", wrap), ul = $("ul", wrap);
  let sel = 0, lista = [];
  const filtra = () => {
    const q = semAcento(inp.value.trim()), dig = inp.value.replace(/\D/g, "");
    lista = st.V.fundos.filter((f) => (!soComCarteira || f.tem_carteira) && (!q || semAcento(f.nome).includes(q) || (dig.length > 2 && f.id.includes(dig))))
      .sort((a, b) => b.nosso - a.nosso || a.nome.localeCompare(b.nome)).slice(0, 30);
    sel = 0;
    ul.innerHTML = lista.map((f, i) => `<li data-i="${i}" class="${i === 0 ? "ativo" : ""}"><span>${f.nosso ? '<span class="ponto" style="background:var(--nosso)"></span>' : ""}${esc(f.nome)}</span><small>${esc(f.cnpj)}</small></li>`).join("") || "<li>Nada encontrado</li>";
    ul.hidden = false;
  };
  const escolhe = (i) => { if (lista[i]) { ul.hidden = true; inp.value = ""; onPick(lista[i]); } };
  inp.addEventListener("input", filtra);
  inp.addEventListener("focus", filtra);
  inp.addEventListener("keydown", (e) => {
    if (e.key === "ArrowDown" || e.key === "ArrowUp") {
      sel = (sel + (e.key === "ArrowDown" ? 1 : -1) + lista.length) % Math.max(lista.length, 1);
      ul.querySelectorAll("li").forEach((x, i) => x.classList.toggle("ativo", i === sel)); e.preventDefault();
    } else if (e.key === "Enter") { e.preventDefault(); escolhe(sel); } else if (e.key === "Escape") ul.hidden = true;
  });
  ul.addEventListener("mousedown", (e) => { const li = e.target.closest("li[data-i]"); if (li) escolhe(+li.dataset.i); });
  inp.addEventListener("blur", () => setTimeout(() => (ul.hidden = true), 150));
  return wrap;
}

// ------------------------------------------------------------------ graficos e tabelas
function layout(extra = {}) {
  return Object.assign({
    paper_bgcolor: "rgba(0,0,0,0)", plot_bgcolor: "rgba(0,0,0,0)", font: { family: css("--font"), size: 12, color: css("--text-2") },
    margin: { l: 50, r: 14, t: 10, b: 40 }, hovermode: "x unified", legend: { orientation: "h", y: -0.18, font: { color: css("--text-2") } },
    xaxis: { gridcolor: "rgba(0,0,0,0)", linecolor: css("--line-2"), automargin: true },
    yaxis: { gridcolor: css("--line"), zerolinecolor: css("--line-2"), automargin: true },
    hoverlabel: { bgcolor: css("--surface"), bordercolor: css("--line-2"), font: { color: css("--text-1") } },
  }, extra);
}
function plota(el, dados, lay, h = 340) {
  el.style.height = h + "px";
  Plotly.newPlot(el, dados, lay, { displaylogo: false, responsive: true, displayModeBar: innerWidth > 700 ? "hover" : false,
    modeBarButtonsToRemove: ["lasso2d", "select2d", "autoScale2d"] });
}
const cartao = (titulo, sub, corpo) => `<div class="cartao"><h3>${titulo}</h3>${sub ? `<p class="sub">${sub}</p>` : ""}${corpo}</div>`;
function tabela(cols, linhas, extra = {}) {
  const th = cols.map((c) => `<th class="${c.n ? "n" : ""} ${c.ord ? "ord" : ""}" ${c.ord ? `data-ord="${c.ord}"` : ""}>${c.t}</th>`).join("");
  const tr = linhas.map((r) => `<tr class="${r._cls || ""}">${cols.map((c) => `<td class="${c.n ? "n" : c.w ? "texto" : ""}">${c.f ? c.f(r[c.k], r) : esc(r[c.k] ?? "–")}</td>`).join("")}</tr>`).join("");
  return `<div class="tabela-wrap" style="max-height:${extra.alt || 620}px"><table><thead><tr>${th}</tr></thead><tbody>${tr}</tbody></table></div>`;
}
const kpi = (r, v, s = "", c = "") => `<div class="kpi"><div class="r">${r}</div><div class="v ${c}">${v}</div><div class="s">${s}</div></div>`;
function categoriasPrincipais(series, n = 7) {
  return Object.entries(series).map(([c, ys]) => [c, ys.reduce((s, y) => s + (y || 0), 0)]).sort((a, b) => b[1] - a[1]).slice(0, n).map((x) => x[0]);
}

function desenha() {
  const el = $("#conteudo");
  try {
    if (st.aba === "carteira") carteira(el);
    else if (st.aba === "retorno") retorno(el);
    else consolidado(el);
  } catch (e) { el.innerHTML = `<div class="vazio">Erro ao desenhar: ${esc(e.message)}</div>`; console.error(e); }
}
function cabecalhoFundo(el, f, extra = "") {
  el.innerHTML = `<div class="cab"><div class="titulo-fundo"><h2>${esc(f.nome)} ${f.nosso ? '<span class="tag">NOSSO</span>' : ""}</h2>
    <p class="sub">${esc(f.cnpj)}${f.pl ? ` · PL R$ ${num(f.pl, 0)} mm` : ""}</p></div><div class="controles" id="ctl">${extra}</div></div><div id="corpo"></div>`;
  $("#ctl").prepend(campoBusca((x) => { st.fundo = x.id; desenha(); }, "Trocar de fundo: digite nome ou CNPJ...", st.aba === "carteira"));
}

// ------------------------------------------------------------------ filtros de ativos (categorias liga/desliga, quantos, busca)
function filtros(box, todas, f, redesenha, placeholder) {
  box.innerHTML = `<div class="filtros"><div class="chips">${todas.map((c) => `<button class="chip-t" data-c="${esc(c)}" aria-pressed="${!f.fora.has(c)}">${esc(c)}</button>`).join("")}
      <button class="chip-t mudo" data-t="1">Todas</button><button class="chip-t mudo" data-t="0">Nenhuma</button></div>
    <div class="controles"><div class="busca"><input type="search" class="q" placeholder="${placeholder}" value="${esc(f.q)}"></div>
      <label>Mostrar <select class="n">${[5, 10, 15, 20, 30, 0].map((n) => `<option value="${n}" ${n === f.n ? "selected" : ""}>${n || "Todos"}</option>`).join("")}</select></label></div></div>`;
  const marca = () => box.querySelectorAll("[data-c]").forEach((x) => x.setAttribute("aria-pressed", !f.fora.has(x.dataset.c)));
  box.querySelectorAll("[data-c]").forEach((b) => (b.onclick = () => { const c = b.dataset.c; f.fora.has(c) ? f.fora.delete(c) : f.fora.add(c); marca(); redesenha(); }));
  box.querySelectorAll("[data-t]").forEach((b) => (b.onclick = () => { f.fora = new Set(b.dataset.t === "1" ? [] : todas); marca(); redesenha(); }));
  $(".q", box).oninput = (e) => { f.q = e.target.value; redesenha(); };
  $(".n", box).onchange = (e) => { f.n = +e.target.value; redesenha(); };
}
const passa = (f, x) => !f.fora.has(x.categoria) && (!f.q.trim() || semAcento([x.ativo, x.grupo, x.emissor, x.codigo].join(" ")).includes(semAcento(f.q.trim())));
const corta = (arr, n) => (n ? arr.slice(0, n) : arr);
const fonteCurta = (v) => `<span class="muted">${esc(v || "–")}</span>`;

// ------------------------------------------------------------------ CARTEIRA
async function carteira(el) {
  const f = fundoPor(st.fundo);
  if (!f.tem_carteira) { cabecalhoFundo(el, f); $("#corpo").innerHTML = `<div class="vazio">A CVM não tem carteira deste fundo no período.</div>`; return; }
  cabecalhoFundo(el, f);
  $("#corpo").innerHTML = `<div class="vazio">Carregando ${esc(f.nome)}...</div>`;
  const D = await detalhe(f.id);
  if (st.fundo !== f.id || st.aba !== "carteira") return;
  const ms = D.carteira.map((x) => x.mes);
  const aberto = [...D.carteira].reverse().find((x) => x.conf < 5) || D.carteira[D.carteira.length - 1];
  const mes = ms.includes(st.mesCart[f.id]) ? st.mesCart[f.id] : aberto.mes;
  const C = D.carteira.find((x) => x.mes === mes);
  $("#ctl").insertAdjacentHTML("beforeend", `<label>Carteira <select id="mes">${ms.slice().reverse().map((m) => `<option value="${m}" ${m === mes ? "selected" : ""}>${mesBR(m)}${D.carteira.find((x) => x.mes === m).conf >= 5 ? " (confid.)" : ""}</option>`).join("")}</select></label>
    <a class="botao" href="/api/csv/${f.id}" title="Carteira look-through completa (CSV)">CSV</a>`);
  $("#mes").onchange = (e) => { st.mesCart[f.id] = e.target.value; desenha(); };

  const linhas = C.cats.map((c) => ({ ...c })).concat([{ ...C.total, _cls: "total", categoria: "Total ponderado" }]);
  const tab = tabela([
    { t: "Categoria", k: "categoria", w: 1 },
    { t: "% PL", k: "perc", n: 1, f: (v) => pct(v) },
    { t: "Financeiro (R$ mm)", k: "fin", n: 1, f: (v) => num(v) },
    { t: "Spread CDI+ (% a.a.)", k: "spread", n: 1, f: (v) => num(v) },
    { t: "Duration (anos)", k: "duration", n: 1, f: (v) => num(v) },
  ], linhas);
  const conf = C.conf >= 0.5 ? ` · ${pct(C.conf, 1)} do PL ainda confidencial` : "";
  const fontes = Object.entries(C.fontes || {}).filter(([, v]) => v >= 0.1).map(([k, v]) => `${esc(k)} ${pct(v, 1)}`).join(" · ");
  $("#corpo").innerHTML = cartao(`Carteira de ${mesBR(mes)} por categoria`, `PL R$ ${num(C.pl_mm, 2)} mm · look-through (cotas de fundos abertas até o ativo)${conf}`, tab +
    `<p class="nota">Spread CDI+ = equivalente em CDI (IPCA+ e prefixados contra a curva do Tesouro de mesma duration). Cobertura de 100% dos ativos de crédito; ações, FII, FIP, ETF e FIAGRO não têm spread. Fonte, em % do PL: ${fontes}.</p>`) +
    `<div class="grade g2">${cartao("% do PL por categoria", "Evolução mensal", '<div id="g_aloc"></div>')}${cartao("Spread CDI+ por categoria", "Média ponderada no mês, % a.a.", '<div id="g_spread"></div>')}</div>
    <div class="cartao"><h3>Ativos de ${mesBR(mes)}</h3><p class="sub">Maiores posições por % do PL · filtre por categoria ou busque por ativo, emissor, grupo econômico ou código</p>
      <div id="f_at"></div><div class="grade g-tabela" style="margin:12px 0 0"><div id="t_top"></div><div><h3>Maiores grupos econômicos</h3><p class="sub">% do PL nas categorias escolhidas</p><div id="g_grupos"></div></div></div></div>`;

  const eixoX = { ...layout().xaxis, tickformat: "%m/%y", dtick: ms.length <= 12 ? "M1" : ms.length <= 36 ? "M3" : "M12" };
  const series = {};
  D.carteira.forEach((x, i) => x.cats.forEach((c) => { (series[c.categoria] ||= Array(D.carteira.length).fill(0))[i] = c.perc; }));
  const tops = categoriasPrincipais(series), cores = st.V.cat_cores;
  const agrup = tops.map((c, i) => ({ c, y: series[c], cor: cores[i % cores.length] }));
  agrup.push({ c: "Outros", y: ms.map((_, i) => Object.entries(series).filter(([c]) => !tops.includes(c)).reduce((s, [, y]) => s + (y[i] || 0), 0)), cor: "#9aa0a6" });
  plota($("#g_aloc"), agrup.map((a) => ({ x: ms.map((m) => m + "-15"), y: a.y, name: a.c, stackgroup: "a", line: { width: 0.5, color: css("--surface") },
    fillcolor: a.cor, hovertemplate: "%{y:.1f}%" })), layout({ xaxis: eixoX, yaxis: { ticksuffix: "%", gridcolor: css("--line") } }), 360);
  const sp = {};
  D.carteira.forEach((x, i) => x.cats.forEach((c) => { if (c.spread != null && c.perc > 0.05) (sp[c.categoria] ||= Array(D.carteira.length).fill(null))[i] = c.spread; }));
  const comSpread = tops.filter((c) => sp[c] && c !== "Caixa" && c !== "Títulos Públicos");
  plota($("#g_spread"), comSpread.map((c) => ({ x: ms.map((m) => m + "-15"), y: sp[c], name: c, mode: "lines", connectgaps: false,
    line: { color: cores[tops.indexOf(c) % cores.length], width: 2 }, hovertemplate: "%{y:.2f}%" })),
    layout({ xaxis: eixoX, yaxis: { ticksuffix: "%", gridcolor: css("--line") } }), 360);

  $("#t_top").innerHTML = `<div class="vazio">Carregando ativos...</div>`;
  const k = f.id + mes;
  if (!st.A[k]) st.A[k] = await (await fetch(`/api/ativos/${f.id}?mes=${mes}`)).json();
  if (st.fundo !== f.id || st.aba !== "carteira" || !$("#t_top")) return;
  const A = st.A[k].map((x) => ({ ...x, ativo: x.rotulo }));
  const peso = {};
  A.forEach((x) => (peso[x.categoria] = (peso[x.categoria] || 0) + x.perc));
  const todas = Object.keys(peso).sort((a, b) => peso[b] - peso[a]);
  const cols = [{ t: "Ativo", k: "ativo", w: 1 }, { t: "Categoria", k: "categoria" }, { t: "Grupo econômico", k: "grupo", w: 1 },
    { t: "% PL", k: "perc", n: 1, f: (v) => pct(v) }, { t: "Spread CDI+", k: "spread", n: 1, f: (v) => num(v) },
    { t: "Duration", k: "duration", n: 1, f: (v) => num(v) }, { t: "PU (R$)", k: "pu", n: 1, f: (v) => num(v) }, { t: "Fonte", k: "fonte", f: fonteCurta }];
  const redesenha = () => {
    const F = st.cf, lista = A.filter((x) => passa(F, x)).sort((a, b) => b.perc - a.perc);
    const soma = lista.reduce((s, x) => s + x.perc, 0);
    $("#t_top").innerHTML = tabela(cols, corta(lista, F.n), { alt: 640 }) +
      `<p class="nota">${lista.length} ativo(s) · ${pct(soma)} do PL nas categorias escolhidas${F.n && lista.length > F.n ? ` · mostrando ${F.n}` : ""}</p>`;
    const gp = {};
    lista.forEach((x) => (gp[x.grupo] = (gp[x.grupo] || 0) + x.perc));
    const g = corta(Object.entries(gp).sort((a, b) => b[1] - a[1]), F.n || 40).reverse();
    plota($("#g_grupos"), [{ type: "bar", orientation: "h", x: g.map((x) => x[1]), y: g.map((x) => x[0]), marker: { color: cores[0] }, hovertemplate: "%{y}: %{x:.2f}% do PL<extra></extra>" }],
      layout({ hovermode: "closest", margin: { l: 8, r: 14, t: 6, b: 30 }, yaxis: { automargin: true, gridcolor: "rgba(0,0,0,0)" }, xaxis: { ticksuffix: "%", gridcolor: css("--line") } }), Math.max(220, 24 * g.length + 50));
  };
  filtros($("#f_at"), todas, st.cf, redesenha, "Buscar ativo, emissor, grupo ou código...");
  redesenha();
}

// ------------------------------------------------------------------ RETORNO
function periodoMeses(ms, chave) {
  if (!ms.includes(st[chave].de)) st[chave].de = ms[Math.max(0, ms.length - 12)];
  if (!ms.includes(st[chave].ate)) st[chave].ate = ms[ms.length - 1];
  const opt = (v) => ms.slice().reverse().map((m) => `<option value="${m}" ${m === v ? "selected" : ""}>${mesBR(m)}</option>`).join("");
  return `<label>De <select id="de">${opt(st[chave].de)}</select></label><label>até <select id="ate">${opt(st[chave].ate)}</select></label>
    ${[["6m", 6], ["12m", 12], ["24m", 24], ["No ano", "ano"], ["Tudo", 999]].map(([t, n]) => `<button class="botao" data-p="${n}">${t}</button>`).join("")}`;
}
function ligaPeriodo(ms, chave) {
  $("#de").onchange = (e) => { st[chave].de = e.target.value; desenha(); };
  $("#ate").onchange = (e) => { st[chave].ate = e.target.value; desenha(); };
  document.querySelectorAll("[data-p]").forEach((b) => (b.onclick = () => {
    const ult = ms[ms.length - 1], n = b.dataset.p;
    st[chave].ate = ult;
    st[chave].de = n === "ano" ? (ms.find((m) => m.startsWith(ult.slice(0, 4))) || ms[0]) : ms[Math.max(0, ms.length - +n)];
    desenha();
  }));
}
function estatPeriodo(f, i0, i1) {
  // retorno, vol anual e drawdown da cota entre os indices i0..i1 de V.datas (precisa existir no inicio e no fim)
  const c = f.cota; if (!c || !c.length) return null;
  let a = i0; while (a <= i1 && c[a] == null) a++;
  if (a > i0 + 6 || a > i1) return null;
  let b = i1; while (b > a && c[b] == null) b--;
  if (b < i1 - 6) return null;
  const rets = []; let pico = c[a], dd = 0, ant = c[a];
  for (let k = a + 1; k <= b; k++) { if (c[k] == null) continue; rets.push(Math.log(c[k] / ant)); ant = c[k]; pico = Math.max(pico, c[k]); dd = Math.min(dd, c[k] / pico - 1); }
  const m = rets.reduce((s, x) => s + x, 0) / Math.max(rets.length, 1);
  const vol = Math.sqrt(rets.reduce((s, x) => s + (x - m) ** 2, 0) / Math.max(rets.length - 1, 1) * 252);
  const cdi = st.V.cdi[b] / st.V.cdi[a] - 1, ret = c[b] / c[a] - 1;
  return { ret: ret * 100, cdi: cdi * 100, pcdi: cdi ? (ret / cdi) * 100 : null, vol: vol * 100, dd: dd * 100 };
}
function indicesDatas(de, ate) {
  const D = st.V.datas;
  let i0 = D.findIndex((d) => d >= de); if (i0 < 0) i0 = D.length - 1;
  let i1 = D.length - 1; while (i1 > 0 && D[i1] > ate) i1--;
  return [Math.max(0, i0 - 1), i1];   // inclui a cota do dia anterior ao inicio (retorno do 1o dia)
}
async function retorno(el) {
  const f = fundoPor(st.fundo);
  cabecalhoFundo(el, f);
  if (!f.tem_carteira) { $("#corpo").innerHTML = `<div class="vazio">A CVM não tem carteira deste fundo no período.</div>`; return; }
  $("#corpo").innerHTML = `<div class="vazio">Carregando ${esc(f.nome)}...</div>`;
  const D = await detalhe(f.id);
  if (st.fundo !== f.id || st.aba !== "retorno") return;
  const R = D.retorno, ms = R.meses.map((x) => x.mes);
  if (!ms.length) { $("#corpo").innerHTML = `<div class="vazio">Sem meses suficientes para calcular retorno.</div>`; return; }
  $("#ctl").insertAdjacentHTML("beforeend", periodoMeses(ms, "ret"));
  ligaPeriodo(ms, "ret");
  const j0 = ms.indexOf(st.ret.de), j1 = ms.indexOf(st.ret.ate);
  const [a, b] = j0 <= j1 ? [j0, j1] : [j1, j0];
  const J = Array.from({ length: b - a + 1 }, (_, k) => a + k);
  const tot = J.reduce((p, j) => p * (1 + R.meses[j].real / 100), 1) - 1, cdi = J.reduce((p, j) => p * (1 + R.meses[j].cdi / 100), 1) - 1;
  const linhas = Object.entries(R.categorias).map(([c, v]) => {
    let contrib = 0, ret = 1, nP = 0, sP = 0, nR = 0;
    J.forEach((j) => { const cc = v.contrib[j] || 0, p = v.peso[j] || 0; contrib += cc; if (p > 0.05) { ret *= 1 + cc / p; nR++; } sP += p; nP++; });
    return { categoria: c, peso: sP / Math.max(nP, 1), contrib, ret: nR ? (ret - 1) * 100 : null, pcdi: nR && cdi ? ((ret - 1) / cdi) * 100 : null };
  }).filter((r) => Math.abs(r.peso) > 0.05 || Math.abs(r.contrib) > 0.005).sort((x, y) => y.contrib - x.contrib);
  linhas.push({ _cls: "total", categoria: "Fundo (retorno da cota)", peso: 100, contrib: tot * 100, ret: tot * 100, pcdi: cdi ? (tot / cdi) * 100 : null });
  const soma = {}, pesoAt = {};
  R.contrib.forEach(([ia, jm, v, w]) => { if (jm >= a && jm <= b) { soma[ia] = (soma[ia] || 0) + v; pesoAt[ia] = (pesoAt[ia] || 0) + (w || 0); } });
  const ativos = Object.keys(soma).map((ia) => ({ ...R.ativos[ia], contrib: soma[ia], perc: pesoAt[ia] / J.length }));
  const hedge = J.reduce((s, j) => s + (R.meses[j].hedge || 0), 0);
  const per = `${mesBR(ms[a])} a ${mesBR(ms[b])}`;

  $("#corpo").innerHTML = `<div class="kpis">${kpi("Retorno do fundo", pct(tot * 100), per, cls(tot))}${kpi("CDI", pct(cdi * 100), "mesmo período")}
      ${kpi("% do CDI", pct(cdi ? (tot / cdi) * 100 : null, 1), "")}${kpi("Excesso sobre o CDI", pp((tot - cdi) * 100), "", cls(tot - cdi))}${kpi("Derivativos", pp(hedge), "já alocado nos ativos protegidos", cls(hedge))}</div>` +
    cartao("Retorno por categoria", "Contribuição em pontos percentuais do PL; retorno e % do CDI da própria categoria no período", tabela([
      { t: "Categoria", k: "categoria", w: 1 }, { t: "Peso médio", k: "peso", n: 1, f: (v) => pct(v, 1) }, { t: "Contribuição", k: "contrib", n: 1, f: (v) => `<span class="${cls(v)}">${pp(v)}</span>` },
      { t: "Retorno", k: "ret", n: 1, f: (v) => pct(v) }, { t: "% do CDI", k: "pcdi", n: 1, f: (v) => pct(v, 0) }], linhas) +
      `<p class="nota">Estimado com as carteiras mensais da CVM: cada ativo rende CDI + seu spread, ou a variação real de preço quando ela é crível; caixa rende CDI; derivativos (DAP → ativos IPCA+, DI1 → prefixados, dólar → ativos no exterior) entram nos ativos que protegem; a diferença para o retorno real da cota (taxas, negociação, marcação) é distribuída pelo peso. A soma das categorias fecha com o retorno real da cota.</p>`) +
    `<div class="cartao"><h3>Ativos por contribuição</h3><p class="sub">${per} · % PL = peso médio no período · filtre por categoria ou busque por ativo, emissor, grupo ou código</p>
      <div id="f_ret"></div><div class="grade g2" style="margin:12px 0 0"><div><h3>Melhores</h3><div id="t_mel"></div></div><div><h3>Piores</h3><div id="t_pio"></div></div></div></div>`;
  const peso = {};
  ativos.forEach((x) => (peso[x.categoria] = (peso[x.categoria] || 0) + Math.abs(x.perc)));
  const todas = Object.keys(peso).sort((x, y) => peso[y] - peso[x]);
  const cols = [{ t: "Ativo", k: "ativo", w: 1 }, { t: "Categoria", k: "categoria" }, { t: "Grupo econômico", k: "grupo", w: 1 },
    { t: "% PL", k: "perc", n: 1, f: (v) => pct(v) }, { t: "Contribuição", k: "contrib", n: 1, f: (v) => `<span class="${cls(v)}">${pp(v, 3)}</span>` }];
  const redesenha = () => {
    const F = st.rf, lista = ativos.filter((x) => passa(F, x));
    $("#t_mel").innerHTML = tabela(cols, corta(lista.slice().sort((x, y) => y.contrib - x.contrib), F.n), { alt: 560 });
    $("#t_pio").innerHTML = tabela(cols, corta(lista.slice().sort((x, y) => x.contrib - y.contrib), F.n), { alt: 560 });
  };
  filtros($("#f_ret"), todas, st.rf, redesenha, "Buscar ativo, emissor, grupo ou código...");
  redesenha();
}
function riscoRetorno(el, linhas) {
  const grupo = (arr, nome, cor) => ({ type: "scatter", mode: "markers+text", name: nome, x: arr.map((r) => r.vol), y: arr.map((r) => r.pcdi),
    text: arr.map((r) => (r.f.nosso || arr.length <= 12 ? r.f.nome : "")), textposition: "top center", textfont: { size: 11, color: css("--text-2") },
    customdata: arr.map((r) => [r.f.nome, r.ret]), marker: { size: 10, color: cor, opacity: 0.9, line: { width: 1, color: css("--surface") } },
    hovertemplate: "%{customdata[0]}<br>retorno %{customdata[1]:.2f}% · %{y:.1f}% do CDI · vol %{x:.2f}%<extra></extra>" });
  plota(el, [grupo(linhas.filter((r) => !r.f.nosso), "Peers", "#9aa0a6"), grupo(linhas.filter((r) => r.f.nosso), "Nossos", "#eb6834")],
    layout({ hovermode: "closest", legend: { orientation: "h", y: -0.2 }, xaxis: { title: "volatilidade anual (%)", gridcolor: css("--line"), zeroline: false },
      yaxis: { title: "% do CDI", gridcolor: css("--line"), zeroline: false } }), 440);
}

// ------------------------------------------------------------------ CONSOLIDADO
function consolidado(el) {
  const V = st.V, C = st.cons, ult = V.datas[V.datas.length - 1];
  if (!C.ate) C.ate = ult;
  if (!C.de) { const d = new Date(ult + "T12:00"); d.setFullYear(d.getFullYear() - 1); C.de = d.toISOString().slice(0, 10); }
  el.innerHTML = `<div class="cab"><div class="titulo-fundo"><h2>Consolidado</h2><p class="sub">Retorno de cada fundo no período · só fundos ativos do início ao fim do período</p></div>
    <div class="controles"><label>De <input type="date" id="cde" value="${C.de}" min="${V.datas[0]}" max="${ult}"></label><label>até <input type="date" id="cate" value="${C.ate}" min="${V.datas[0]}" max="${ult}"></label>
    ${[["12m", 12], ["24m", 24], ["No ano", "ano"], ["Tudo", "tudo"]].map(([t, n]) => `<button class="botao" data-q="${n}">${t}</button>`).join("")}</div></div>
    <div class="cartao"><div class="controles" style="margin-bottom:8px"><span class="sub" style="margin:0">Fundos:</span>
      <button class="botao" data-s="todos">Todos</button><button class="botao" data-s="nossos">Nossos</button><button class="botao" data-s="peers">Peers</button><button class="botao" data-s="nenhum">Limpar</button></div>
      <div class="multi" id="multi"></div></div>
    <div id="cc"></div>`;
  $("#cde").onchange = (e) => { C.de = e.target.value; desenha(); };
  $("#cate").onchange = (e) => { C.ate = e.target.value; desenha(); };
  document.querySelectorAll("[data-q]").forEach((b) => (b.onclick = () => {
    const d = new Date(ult + "T12:00"), q = b.dataset.q;
    C.ate = ult;
    C.de = q === "tudo" ? V.datas[0] : q === "ano" ? ult.slice(0, 4) + "-01-01" : (d.setMonth(d.getMonth() - +q), d.toISOString().slice(0, 10));
    desenha();
  }));
  document.querySelectorAll("[data-s]").forEach((b) => (b.onclick = () => {
    const s = b.dataset.s;
    C.sel = new Set(V.fundos.filter((f) => s === "todos" || (s === "nossos" && f.nosso) || (s === "peers" && !f.nosso)).map((f) => f.id));
    desenha();
  }));
  const multi = $("#multi");
  const chips = V.fundos.filter((f) => C.sel.has(f.id));
  multi.innerHTML = (chips.length > 24 ? `<span class="chip">${chips.filter((f) => f.nosso).length} nossos + ${chips.filter((f) => !f.nosso).length} peers selecionados</span>`
    : chips.map((f) => `<span class="chip ${f.nosso ? "nosso" : ""}">${esc(f.nome)}<button data-x="${f.id}" aria-label="Remover">×</button></span>`).join(""));
  multi.querySelectorAll("[data-x]").forEach((b) => (b.onclick = () => { C.sel.delete(b.dataset.x); desenha(); }));
  multi.appendChild(campoBusca((x) => { C.sel.add(x.id); desenha(); }, "Adicionar fundo: digite..."));

  const [i0, i1] = indicesDatas(C.de, C.ate);
  const linhas = [], fora = [];
  V.fundos.filter((f) => C.sel.has(f.id)).forEach((f) => { const e = estatPeriodo(f, i0, i1); (e ? linhas : fora).push({ f, ...(e || {}) }); });
  const cdiP = (V.cdi[i1] / V.cdi[i0] - 1) * 100;
  const ord = C.ord, sgn = C.asc ? 1 : -1;
  linhas.sort((x, y) => sgn * ((x[ord] ?? -1e9) - (y[ord] ?? -1e9)));
  $("#cc").innerHTML = cartao("Retorno acumulado", `${dataBR(V.datas[i0 + 1] || V.datas[i0])} a ${dataBR(V.datas[i1])} · CDI ${pct(cdiP)} · laranja = nossos · cinza = peers`, '<div id="g_cons"></div>') +
    cartao("Risco x retorno", "Mesmo período e mesmos fundos · laranja = nossos · cinza = peers", '<div id="g_rr"></div>') +
    cartao("Ranking no período", fora.length ? `${fora.length} fundo(s) selecionado(s) fora por não existirem no período inteiro: ${fora.map((x) => esc(x.f.nome)).join(", ")}` : "Clique no título da coluna para ordenar",
      tabela([{ t: "#", k: "_i" }, { t: "Fundo", k: "nome", w: 1, f: (v, r) => `${r.f.nosso ? '<span class="ponto" style="background:var(--nosso)"></span>' : '<span class="ponto" style="background:#9aa0a6"></span>'}${esc(r.f.nome)}` },
        { t: "Retorno", k: "ret", n: 1, ord: "ret", f: (v) => `<span class="${cls(v)}">${pct(v)}</span>` }, { t: "% do CDI", k: "pcdi", n: 1, ord: "pcdi", f: (v) => pct(v, 1) },
        { t: "Vol. anual", k: "vol", n: 1, ord: "vol", f: (v) => pct(v) }, { t: "Pior drawdown", k: "dd", n: 1, ord: "dd", f: (v) => pct(v) },
        { t: "PL (R$ mm)", k: "pl", n: 1, ord: "pl", f: (v) => num(v, 0) }],
        linhas.map((r, i) => ({ ...r, _i: i + 1, nome: r.f.nome, pl: r.f.pl, _cls: r.f.nosso ? "nosso" : "" }))));
  document.querySelectorAll("th[data-ord]").forEach((th) => (th.onclick = () => { if (C.ord === th.dataset.ord) C.asc = !C.asc; else { C.ord = th.dataset.ord; C.asc = false; } desenha(); }));
  const xs = V.datas.slice(i0, i1 + 1);
  const serie = (f) => { const c = f.cota.slice(i0, i1 + 1); let b = c.find((v) => v != null); return c.map((v) => (v == null ? null : (v / b - 1) * 100)); };
  const dados = linhas.filter((r) => !r.f.nosso).map((r) => ({ x: xs, y: serie(r.f), name: r.f.nome, mode: "lines", line: { color: "#9aa0a6", width: 1.2 }, opacity: 0.55, showlegend: false,
    hovertemplate: `${esc(r.f.nome)}: %{y:.2f}%<extra></extra>` }));
  linhas.filter((r) => r.f.nosso).forEach((r) => dados.push({ x: xs, y: serie(r.f), name: r.f.nome, mode: "lines", line: { color: r.f.cor, width: 2.8 }, hovertemplate: `${esc(r.f.nome)}: %{y:.2f}%<extra></extra>` }));
  const c0 = V.cdi[i0];
  dados.push({ x: xs, y: V.cdi.slice(i0, i1 + 1).map((v) => (v / c0 - 1) * 100), name: "CDI", mode: "lines", line: { color: css("--text-1"), width: 2, dash: "dash" }, hovertemplate: "CDI: %{y:.2f}%<extra></extra>" });
  plota($("#g_cons"), dados, layout({ hovermode: "closest", yaxis: { ticksuffix: "%", gridcolor: css("--line") } }), 460);
  riscoRetorno($("#g_rr"), linhas);
}

</script>
</body>
</html>
"""


def gera_excel():
    CNPJ = {n: cnpj_of(x) for n, x in {**NOSSOS_FUNDOS, **PEERS}.items()}
    meses, dia = prepara(sorted(CNPJ.values()), DESDE)
    print(f"{len(meses)} meses de carteira ({meses[0]} a {meses[-1]})")
    df = lookthrough(CNPJ, meses)
    soma = df[~df.derivativo].groupby(["fundo", "mes"]).perc_pl.sum()
    print("soma do %PL por carteira (sem derivativos): min", round(soma.min(), 3), "max", round(soma.max(), 3))
    grandes = df[df.categoria.eq(AJUSTE) & (df.perc_pl.abs() > 1)]
    if len(grandes):
        print("ajustes carteira x PL acima de 1% do peer (veiculo onde a CDA nao fecha com o PL):")
        print(grandes.groupby(["fundo", "veiculo"]).perc_pl.agg(["count", "min", "max"]).round(2).to_string())
    cruz = df[df.cnpj_veiculo.isin(CNPJ.values()) & (df.cnpj_veiculo != df.fundo.map(CNPJ))]
    print("peers investindo em outro peer:", "nenhum" if cruz.empty else sorted(set(zip(cruz.fundo, cruz.veiculo))))
    rent = rentabilidade(diario(CNPJ, dia), cdi_mensal(DESDE))
    excel(df, rent, list(CNPJ), OUT)
    print("salvo", OUT, "e", OUT.replace(".xlsx", "_dados.csv"), len(df), "linhas")


if __name__ == "__main__":                 # importante: os processos auxiliares do Windows reimportam este arquivo
    if len(sys.argv) > 1 and sys.argv[1].lower() == "excel":
        gera_excel()
    else:
        import webbrowser
        import uvicorn
        print(f"Preparando o painel: {len(NOSSOS_FUNDOS) + len(PEERS)} fundos desde {DESDE}. Na 1a vez pode levar bastante tempo...")
        precarrega(log=lambda m: print("  " + m, flush=True))
        if PRE["status"] != "ok":
            sys.exit("Nao foi possivel preparar o painel: " + str(PRE["erro"]))
        porta = int(os.environ.get("PORT", 7860))
        threading.Timer(1.5, lambda: webbrowser.open(f"http://localhost:{porta}")).start()
        print(f"Pronto. Painel em http://localhost:{porta}  (deixe esta janela aberta; Ctrl+C para fechar)")
        uvicorn.run(app, host="127.0.0.1", port=porta, log_level="warning")
